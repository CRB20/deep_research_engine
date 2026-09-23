#!/usr/bin/env python3
"""Interactive follow-up research assistant for completed Deep Research Engine runs.

Features
--------
* Maintains/refreshes runs/research_registry.json from completed run directories.
* Shows the latest 10 research topics, their date/age, pilot vs expanded status,
  completion status, and whether a refresh is sensible.
* Lets the user expand to older topics without loading the entire registry into
  the interactive display.
* Uses the selected run's local evidence corpus first.
* By default performs a fresh external search before answering, with optional:
  Google (via SerpApi when configured), web-search fallback, YouTube, GitHub,
  Reddit, Quora, OpenAlex, arXiv, Crossref, and DBLP.
* Every external source is slow/fail-soft and can be disabled for the current
  question after repeated failures.
* Saves a complete follow-up trace under followup_runs/<timestamp>/.

This script deliberately does not modify the original research-run corpus.
"""
from __future__ import annotations

import argparse
import base64
import concurrent.futures
import hashlib
import math
import shutil
import subprocess
import json
import logging
import mimetypes
import os
import re
import sqlite3
import sys
import threading
import time
import xml.etree.ElementTree as ET
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import yaml
from urllib.parse import quote_plus

import requests
from dotenv import load_dotenv

try:
    from ddgs import DDGS
except Exception:
    DDGS = None

try:
    from rapidfuzz.fuzz import ratio as fuzz_ratio
except Exception:
    def fuzz_ratio(a: str, b: str) -> float:
        from difflib import SequenceMatcher
        return SequenceMatcher(None, a, b).ratio() * 100.0

try:
    from langchain_ollama import ChatOllama
except Exception:
    ChatOllama = None

ROOT = Path(__file__).resolve().parent
ASSISTANT_ENV_PATH = ROOT / os.getenv("RESEARCH_ASSISTANT_ENV_FILE", ".env.research_assistant")
load_dotenv(ASSISTANT_ENV_PATH, override=True)

RUNS_DIR = ROOT / "runs"
FOLLOWUP_RUNS_DIR = ROOT / "followup_runs"
QUICK_RESEARCH_DIR = ROOT / "quick_research_runs"
CONVERSATION_SESSIONS_DIR = ROOT / "conversation_sessions"
DOCUMENT_CONFIG_PATH = ROOT / "document_library.yaml"
RAG_DIR = ROOT / "rag_index"
RAG_DB_PATH = RAG_DIR / "documents.sqlite"
REGISTRY_PATH = RUNS_DIR / "research_registry.json"
IMAGE_DIR = ROOT / "images"
FOLLOWUP_RUNS_DIR.mkdir(parents=True, exist_ok=True)
QUICK_RESEARCH_DIR.mkdir(parents=True, exist_ok=True)
CONVERSATION_SESSIONS_DIR.mkdir(parents=True, exist_ok=True)
RAG_DIR.mkdir(parents=True, exist_ok=True)
IMAGE_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Three broad user-intention categories
# ---------------------------------------------------------------------------
INTENT_CHAT = "CHAT"
INTENT_RESEARCH = "RESEARCH"
INTENT_DOCUMENTS = "DOCUMENT_UNDERSTANDING"
INTENTS = (INTENT_CHAT, INTENT_RESEARCH, INTENT_DOCUMENTS)

def env_bool(name: str, default: bool) -> bool:
    return os.getenv(name, "true" if default else "false").strip().lower() in {"1", "true", "yes", "y", "on"}

# Keep third-party PDF parser diagnostics out of normal interactive output.
# Detailed diagnostics remain available through RAG_VERBOSE_LOGGING.
logging.getLogger("pypdf").setLevel(logging.ERROR)
logging.getLogger("pypdf._font").setLevel(logging.ERROR)
logging.getLogger("pypdf._cmap").setLevel(logging.ERROR)

# Startup / category configuration. The assistant always asks for intention by default.
ASK_INTENTION_ON_START = env_bool("ASSISTANT_ASK_INTENTION_ON_START", True)
CHAT_ENABLED = env_bool("CHAT_CATEGORY_ENABLED", True)
RESEARCH_ENABLED = env_bool("RESEARCH_CATEGORY_ENABLED", True)
DOCUMENTS_ENABLED = env_bool("DOCUMENTS_CATEGORY_ENABLED", True)

CHAT_ALLOW_WEB_SEARCH = env_bool("CHAT_ALLOW_WEB_SEARCH", True)
CHAT_DEFAULT_WEB_SEARCH = env_bool("CHAT_DEFAULT_WEB_SEARCH", False)

RESEARCH_ALLOW_DATABASE = env_bool("RESEARCH_ALLOW_DATABASE", True)
RESEARCH_DEFAULT_DATABASE = env_bool("RESEARCH_DEFAULT_DATABASE", True)
RESEARCH_ALLOW_WEB_SEARCH = env_bool("RESEARCH_ALLOW_WEB_SEARCH", True)
RESEARCH_DEFAULT_WEB_SEARCH = env_bool("RESEARCH_DEFAULT_WEB_SEARCH", True)
RESEARCH_DATABASE_MAX_PROJECTS = int(os.getenv("RESEARCH_DATABASE_MAX_PROJECTS", "20"))

DOCUMENTS_ALLOW_WEB_SEARCH = env_bool("DOCUMENTS_ALLOW_WEB_SEARCH", True)
DOCUMENTS_DEFAULT_WEB_SEARCH = env_bool("DOCUMENTS_DEFAULT_WEB_SEARCH", False)
DOCUMENTS_REQUIRE_PROJECT = env_bool("DOCUMENTS_REQUIRE_PROJECT", True)

# ---------------- Deep Research Engine integration ----------------
# The assistant and Deep Research Engine intentionally share the same workspace.
# The assistant can launch the engine and then resume the same conversation.
DEEP_RESEARCH_ENGINE_SCRIPT = os.getenv("DEEP_RESEARCH_ENGINE_SCRIPT", "deep_research_engine.py")
DEEP_RESEARCH_ENGINE_MODE = os.getenv("DEEP_RESEARCH_ENGINE_MODE", "deep").strip().lower()
if DEEP_RESEARCH_ENGINE_MODE not in {"short", "deep"}:
    DEEP_RESEARCH_ENGINE_MODE = "deep"
DEEP_RESEARCH_AUTO_RESUME = env_bool("DEEP_RESEARCH_AUTO_RESUME", True)
DEEP_RESEARCH_PASS_MODE = env_bool("DEEP_RESEARCH_PASS_MODE", True)
DEEP_RESEARCH_WORKING_DIR = os.getenv("DEEP_RESEARCH_WORKING_DIR", "").strip()

# Legacy compatibility: old --mode values are mapped to the new three-category model.
LEGACY_MODES = ("AUTO", "DOCUMENTS_ONLY", "RESEARCH_ONLY", "DOCUMENTS + RESEARCH", "WEB_SEARCH")
MODE_AUTO = "AUTO"
MODE_DOCUMENTS_ONLY = "DOCUMENTS_ONLY"
MODE_RESEARCH_ONLY = "RESEARCH_ONLY"
MODE_DOCUMENTS_RESEARCH = "DOCUMENTS + RESEARCH"
MODE_WEB_SEARCH = "WEB_SEARCH"

OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434").rstrip("/")
RAG_EMBEDDINGS_ENABLED = os.getenv("RAG_EMBEDDINGS_ENABLED", "false").lower() == "true"
RAG_EMBEDDING_MODEL = os.getenv("RAG_EMBEDDING_MODEL", "nomic-embed-text")
RAG_TOP_K = int(os.getenv("RAG_TOP_K", "10"))
RAG_CANDIDATE_K = int(os.getenv("RAG_CANDIDATE_K", "40"))
RAG_CHUNK_SIZE = int(os.getenv("RAG_CHUNK_SIZE", "1800"))
RAG_CHUNK_OVERLAP = int(os.getenv("RAG_CHUNK_OVERLAP", "250"))
RAG_MAX_FILE_MB = int(os.getenv("RAG_MAX_FILE_MB", "80"))
RAG_INDEX_ON_START = os.getenv("RAG_INDEX_ON_START", "false").lower() == "true"
RAG_SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
RAG_OCR_ENABLED = env_bool("RAG_OCR_ENABLED", True)
RAG_OCR_DPI = int(os.getenv("RAG_OCR_DPI", "180"))
RAG_OCR_LANG = os.getenv("RAG_OCR_LANG", "eng").strip() or "eng"
RAG_OCR_MIN_PAGE_CHARS = int(os.getenv("RAG_OCR_MIN_PAGE_CHARS", "80"))
RAG_OCR_MIN_NATIVE_COVERAGE = float(os.getenv("RAG_OCR_MIN_NATIVE_COVERAGE", "0.65"))
RAG_OCR_MAX_PAGES = int(os.getenv("RAG_OCR_MAX_PAGES", "0"))
RAG_OCR_TIMEOUT_SECONDS = int(os.getenv("RAG_OCR_TIMEOUT_SECONDS", "120"))
RAG_VERBOSE_LOGGING = env_bool("RAG_VERBOSE_LOGGING", False)

# ------------------------- Multimodal image input -------------------------
# Dedicated visual model: analyze attached images first, then pass the visual
# observations to the main answer model.
VISION_ENABLED = env_bool("VISION_ENABLED", True)
VISION_MODEL = os.getenv("VISION_MODEL", "qwen3.8:27b").strip() or "qwen3.8:27b"
VISION_CTX = int(os.getenv("VISION_CTX", "12288"))
VISION_TOKENS = int(os.getenv("VISION_TOKENS", "3000"))
VISION_TIMEOUT_SECONDS = float(os.getenv("VISION_TIMEOUT_SECONDS", "1200"))
VISION_MAX_IMAGE_MB = int(os.getenv("VISION_MAX_IMAGE_MB", "20"))
VISION_MAX_IMAGES = int(os.getenv("VISION_MAX_IMAGES", "4"))
VISION_SERIAL = env_bool("VISION_SERIAL", True)
VISION_SAVE_EVIDENCE = env_bool("VISION_SAVE_EVIDENCE", True)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

QUERY_MODEL = os.getenv("FOLLOWUP_QUERY_MODEL", "qwen3.5:35b-a3b")
ANSWER_MODEL = os.getenv("FOLLOWUP_ANSWER_MODEL", "qwen3.5:35b-a3b")
QUERY_CTX = int(os.getenv("FOLLOWUP_QUERY_CTX", "6144"))
QUERY_TOKENS = int(os.getenv("FOLLOWUP_QUERY_TOKENS", "700"))
ANSWER_CTX = int(os.getenv("FOLLOWUP_ANSWER_CTX", "12288"))
ANSWER_TOKENS = int(os.getenv("FOLLOWUP_ANSWER_TOKENS", "2500"))
ANSWER_HEARTBEAT = float(os.getenv("FOLLOWUP_LLM_HEARTBEAT", "15"))
LLM_TIMEOUT = float(os.getenv("FOLLOWUP_LLM_TIMEOUT_SECONDS", "1200"))
SUMMARY_MODEL = os.getenv("FOLLOWUP_SUMMARY_MODEL", QUERY_MODEL)
SUMMARY_CTX = int(os.getenv("FOLLOWUP_SUMMARY_CTX", "8192"))
SUMMARY_TOKENS = int(os.getenv("FOLLOWUP_SUMMARY_TOKENS", "1800"))
CONVERSATION_SUMMARY_CHARS = int(os.getenv("FOLLOWUP_CONVERSATION_SUMMARY_CHARS", "6000"))
CONVERSATION_HISTORY_CHARS = int(os.getenv("FOLLOWUP_CONVERSATION_HISTORY_CHARS", "12000"))
CONVERSATION_RECENT_TURNS = int(os.getenv("FOLLOWUP_CONVERSATION_RECENT_TURNS", "4"))
CONVERSATION_SUMMARIZE_AFTER_TURNS = int(os.getenv("FOLLOWUP_CONVERSATION_SUMMARIZE_AFTER_TURNS", "4"))
QUICK_QUERY_MODEL = os.getenv("QUICK_QUERY_MODEL", QUERY_MODEL)
QUICK_ANSWER_MODEL = os.getenv("QUICK_ANSWER_MODEL", ANSWER_MODEL)
QUICK_QUERY_CTX = int(os.getenv("QUICK_QUERY_CTX", "8192"))
QUICK_QUERY_TOKENS = int(os.getenv("QUICK_QUERY_TOKENS", "1000"))
QUICK_ANSWER_CTX = int(os.getenv("QUICK_ANSWER_CTX", "12288"))
QUICK_ANSWER_TOKENS = int(os.getenv("QUICK_ANSWER_TOKENS", "2500"))
QUICK_RESULTS_PER_SOURCE = int(os.getenv("QUICK_RESULTS_PER_SOURCE", "4"))
QUICK_MAX_WORKERS = int(os.getenv("QUICK_MAX_WORKERS", "4"))
QUICK_MAX_SOURCE_FAMILIES = int(os.getenv("QUICK_MAX_SOURCE_FAMILIES", "6"))
QUICK_CONTEXT_CHARS = int(os.getenv("QUICK_CONTEXT_CHARS", "30000"))

ALWAYS_REFRESH = os.getenv("FOLLOWUP_ALWAYS_REFRESH", "true").lower() == "true"
RECENT_DAYS = float(os.getenv("FOLLOWUP_RECENT_DAYS", "7"))
PAGE_SIZE = int(os.getenv("FOLLOWUP_REGISTRY_PAGE_SIZE", "10"))
LOCAL_PAPER_RESULTS = int(os.getenv("FOLLOWUP_LOCAL_PAPER_RESULTS", "6"))
LOCAL_EVIDENCE_RESULTS = int(os.getenv("FOLLOWUP_LOCAL_EVIDENCE_RESULTS", "8"))
LOCAL_TECH_RESULTS = int(os.getenv("FOLLOWUP_LOCAL_TECH_RESULTS", "6"))
EXTERNAL_RESULTS_PER_SOURCE = int(os.getenv("FOLLOWUP_EXTERNAL_RESULTS_PER_SOURCE", "3"))
EXTERNAL_MAX_WORKERS = int(os.getenv("FOLLOWUP_EXTERNAL_MAX_WORKERS", "4"))
FOLLOWUP_QUERIES_PER_FAMILY = int(os.getenv("FOLLOWUP_QUERIES_PER_FAMILY", "1"))
FOLLOWUP_CONTEXT_CHARS = int(os.getenv("FOLLOWUP_CONTEXT_CHARS", "30000"))

SERPAPI_API_KEY = os.getenv("SERPAPI_API_KEY", "").strip()
YOUTUBE_API_KEY = os.getenv("YOUTUBE_API_KEY", "").strip()
GITHUB_TOKEN = os.getenv("GITHUB_TOKEN", "").strip()

OPENALEX_API_KEY = os.getenv("OPENALEX_API_KEY", "").strip()
SEMANTIC_SCHOLAR_API_KEY = os.getenv("SEMANTIC_SCHOLAR_API_KEY", "").strip()
SEMANTIC_SCHOLAR_ENABLED = os.getenv("SEMANTIC_SCHOLAR_ENABLED", "true").lower() == "true"

UA = os.getenv(
    "RESEARCH_USER_AGENT",
    "LocalResearchAssistant/1.0 (personal academic research tool)",
)
HEADERS = {"User-Agent": UA}

SOURCE_CONFIG = {
    "google": {"enabled": True, "interval": 3.0, "failures": 2},
    "web": {"enabled": DDGS is not None, "interval": 2.0, "failures": 2},
    "youtube": {"enabled": True, "interval": 3.0, "failures": 2},
    "github": {"enabled": True, "interval": 4.0, "failures": 2},
    "reddit": {"enabled": DDGS is not None, "interval": 4.0, "failures": 2},
    "quora": {"enabled": DDGS is not None, "interval": 4.0, "failures": 2},
    "openalex": {"enabled": True, "interval": 2.0, "failures": 2},
    "semantic_scholar": {"enabled": SEMANTIC_SCHOLAR_ENABLED, "interval": 2.0, "failures": 2},
    "arxiv": {"enabled": True, "interval": 3.0, "failures": 2},
    "crossref": {"enabled": True, "interval": 1.5, "failures": 2},
    "dblp": {"enabled": True, "interval": 2.0, "failures": 2},
}

SOURCE_LOCKS = {name: threading.Lock() for name in SOURCE_CONFIG}
SOURCE_LAST = {name: 0.0 for name in SOURCE_CONFIG}
SOURCE_ERRORS = {name: 0 for name in SOURCE_CONFIG}
SOURCE_DISABLED = {name: False for name in SOURCE_CONFIG}
SOURCE_FAILURE_REASONS: dict[str, str] = {}
LAST_EXTERNAL_STATUS: dict[str, Any] = {
    "attempted": False,
    "failed": False,
    "elapsed": 0.0,
    "results": 0,
    "source_counts": {},
    "failures": {},
}

# ANSI terminal colours. Automatically disabled for non-interactive output or NO_COLOR.
USE_COLOR = bool(sys.stdout.isatty()) and not os.getenv("NO_COLOR")
ANSI = {
    "reset": "\033[0m",
    "dim": "\033[2m",
    "cyan": "\033[96m",
    "green": "\033[92m",
    "yellow": "\033[93m",
    "red": "\033[91m",
    "blue": "\033[94m",
    "magenta": "\033[95m",
}

def colour(text: str, name: str) -> str:
    if not USE_COLOR:
        return text
    return f"{ANSI.get(name, '')}{text}{ANSI['reset']}"

def terminal_log(tag: str, message: str, color: str = "dim") -> None:
    print(colour(f"[{tag}]", color) + f" {message}", flush=True)


@dataclass
class Result:
    source: str
    title: str
    url: str
    snippet: str = ""
    date: str = ""
    metadata: dict[str, Any] | None = None

    def __post_init__(self):
        if self.metadata is None:
            self.metadata = {}


@dataclass
class ConversationState:
    topic: str
    intention: str = INTENT_CHAT
    research_run: dict[str, Any] | None = None
    document_project: str | None = None
    chat_web_search: bool = False
    research_use_database: bool = True
    research_web_search: bool = True
    documents_web_search: bool = False
    summary: str = ""
    turns: list[dict[str, str]] | None = None
    session_id: str = ""
    # Intention/project-scoped memories. Each scope stores its own summary + turns.
    # Examples:
    #   "chat"
    #   "research:/path/to/runs/..."
    #   "research:new:<topic>"
    #   "documents:ROV_Control"
    memory_scopes: dict[str, dict[str, Any]] | None = None
    # Persistent image attachments for the active conversation/topic.
    attached_images: list[str] | None = None

    # Backward-compatible fields for old JSON sessions/CLI options.
    mode: str = ""
    project: dict[str, Any] | None = None

    def __post_init__(self):
        if self.turns is None:
            self.turns = []
        if self.memory_scopes is None:
            self.memory_scopes = {}
        if self.attached_images is None:
            self.attached_images = []
        if self.research_run is None and self.project is not None:
            self.research_run = self.project
        if self.project is None and self.research_run is not None:
            self.project = self.research_run
        if self.intention not in INTENTS:
            legacy = str(self.mode or "").upper()
            if legacy == "DOCUMENTS_ONLY":
                self.intention = INTENT_DOCUMENTS
                self.documents_web_search = False
            elif legacy in {"RESEARCH_ONLY", "DOCUMENTS + RESEARCH", "AUTO"}:
                self.intention = INTENT_RESEARCH
                self.research_use_database = bool(self.research_run)
                self.research_web_search = legacy != "RESEARCH_ONLY" or True
            elif legacy == "WEB_SEARCH":
                self.intention = INTENT_CHAT
                self.chat_web_search = True
            else:
                self.intention = INTENT_CHAT
        self.sync_legacy_mode()

    def sync_legacy_mode(self) -> None:
        if self.intention == INTENT_CHAT:
            self.mode = "CHAT + WEB" if self.chat_web_search else "CHAT"
        elif self.intention == INTENT_RESEARCH:
            parts = []
            if self.research_use_database:
                parts.append("DATABASE")
            if self.research_web_search:
                parts.append("WEB")
            self.mode = "RESEARCH" + (" [" + " + ".join(parts) + "]" if parts else " [LLM ONLY]")
        else:
            self.mode = "DOCUMENTS + WEB" if self.documents_web_search else "DOCUMENTS_ONLY"



def safe_slug(text: str, max_len: int = 60) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", str(text).strip()).strip("._-")
    return (value[:max_len] or "topic")


def memory_scope_key(state: ConversationState) -> str:
    """Return the persistent memory namespace for the currently active context."""
    if state.intention == INTENT_CHAT:
        return "chat"
    if state.intention == INTENT_RESEARCH:
        if state.research_run and state.research_run.get("run_dir"):
            return f"research:{state.research_run['run_dir']}"
        topic = safe_slug(state.topic or "new-research", max_len=100)
        return f"research:new:{topic}"
    if state.intention == INTENT_DOCUMENTS:
        return f"documents:{state.document_project or 'none'}"
    return f"other:{safe_slug(state.topic or 'session')}"


def sync_active_memory_scope(state: ConversationState) -> None:
    """Persist the active summary/turns into its intention-specific memory scope."""
    if state.memory_scopes is None:
        state.memory_scopes = {}
    key = memory_scope_key(state)
    state.memory_scopes[key] = {
        "summary": state.summary or "",
        "turns": list(state.turns or []),
    }


def activate_memory_scope(state: ConversationState, *, preserve_current: bool = True) -> None:
    """Switch the visible conversation memory to the active intention/project scope."""
    if preserve_current:
        sync_active_memory_scope(state)
    if state.memory_scopes is None:
        state.memory_scopes = {}
    key = memory_scope_key(state)
    bucket = state.memory_scopes.get(key) or {}
    state.summary = str(bucket.get("summary") or "")
    turns = bucket.get("turns") or []
    state.turns = list(turns) if isinstance(turns, list) else []


def conversation_prompt_context(state: ConversationState) -> str:
    """Return a bounded conversational-memory block while preserving the summary."""
    summary_budget = min(CONVERSATION_SUMMARY_CHARS, CONVERSATION_HISTORY_CHARS // 2)
    recent_budget = max(1000, CONVERSATION_HISTORY_CHARS - summary_budget)
    parts: list[str] = []
    if state.summary.strip():
        parts.append("CONVERSATION SUMMARY (compressed memory; not a source of factual evidence):\n" +
                     state.summary.strip()[:summary_budget])
    recent = state.turns[-CONVERSATION_RECENT_TURNS:] if state.turns else []
    if recent:
        turns = []
        for i, turn in enumerate(recent, 1):
            turns.append(
                f"TURN {i} USER:\n{turn.get('question','')[:3500]}\n\n"
                f"TURN {i} ASSISTANT:\n{turn.get('answer','')[:5500]}"
            )
        recent_text = "RECENT CONVERSATION TURNS:\n" + "\n\n".join(turns)
        parts.append(recent_text[-recent_budget:])
    return "\n\n".join(parts)


def add_conversation_turn(state: ConversationState, question: str, answer: str) -> None:
    state.turns.append({"question": question, "answer": answer})
    # Keep a generous in-memory window; the summary compresses older material.
    if len(state.turns) > 12:
        state.turns = state.turns[-12:]


def refresh_conversation_summary(state: ConversationState) -> None:
    """Compress older conversation into a persistent summary only when useful."""
    if not state.turns or len(state.turns) < CONVERSATION_SUMMARIZE_AFTER_TURNS:
        return
    if len(state.turns) <= CONVERSATION_RECENT_TURNS:
        return
    older = state.turns[:-CONVERSATION_RECENT_TURNS]
    prior = state.summary.strip() or "(no previous summary)"
    transcript = "\n\n".join(
        f"USER: {t.get('question','')}\nASSISTANT: {t.get('answer','')}" for t in older
    )
    prompt = f"""
Maintain a compact research-conversation memory.

Existing summary:
{prior[:CONVERSATION_SUMMARY_CHARS]}

Older conversation turns to compress:
{transcript[:50000]}

Create a concise factual memory containing:
- the user's research goals and constraints
- questions already answered
- important conclusions or hypotheses discussed
- unresolved questions / uncertainties
- named papers, methods, systems, terminology, or entities that matter for continuity
- decisions the user made
Do NOT invent facts and do NOT treat conversation claims as verified evidence.
Do not include greetings or filler.
Return only the memory summary.
"""
    state.summary = invoke_model(
        SUMMARY_MODEL, prompt, ctx=SUMMARY_CTX,
        tokens_out=SUMMARY_TOKENS, heartbeat=ANSWER_HEARTBEAT,
    )[:CONVERSATION_SUMMARY_CHARS]
    state.turns = state.turns[-CONVERSATION_RECENT_TURNS:]


def save_conversation_state(state: ConversationState) -> Path:
    # Keep the persisted scope synchronized with the currently visible memory.
    sync_active_memory_scope(state)
    session_id = state.session_id or datetime.now().strftime("%Y%m%d_%H%M%S")
    state.session_id = session_id
    out = CONVERSATION_SESSIONS_DIR / f"{session_id}_{safe_slug(state.topic)}.json"
    payload = asdict(state)
    payload["updated"] = datetime.now(timezone.utc).isoformat()
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def load_latest_conversation_state(topic: str, research_run: dict[str, Any] | None,
                                     document_project: str | None = None) -> ConversationState | None:
    """Resume the most recent saved conversation for a research run/topic."""
    candidates = []
    for path in CONVERSATION_SESSIONS_DIR.glob("*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        saved_run = data.get("research_run") or data.get("project")
        same = False
        if research_run and saved_run:
            same = bool(research_run.get("run_dir") and saved_run.get("run_dir") == research_run.get("run_dir"))
        elif not research_run:
            same = (data.get("topic") == topic and data.get("document_project") == document_project)
        if same:
            candidates.append((path.stat().st_mtime, data))
    if not candidates:
        return None
    _, data = max(candidates, key=lambda x: x[0])
    state = ConversationState(
        topic=data.get("topic", topic),
        intention=data.get("intention", ""),
        research_run=data.get("research_run") or data.get("project"),
        document_project=data.get("document_project"),
        chat_web_search=bool(data.get("chat_web_search", False)),
        research_use_database=bool(data.get("research_use_database", True)),
        research_web_search=bool(data.get("research_web_search", True)),
        documents_web_search=bool(data.get("documents_web_search", False)),
        summary=data.get("summary", ""),
        turns=data.get("turns", []),
        session_id=data.get("session_id", ""),
        memory_scopes=data.get("memory_scopes") or {},
        mode=data.get("mode", ""),
        project=data.get("project"),
    )
    # Migrate older single-buffer sessions: assign the legacy turns to the
    # context that was active when the old session was saved.
    if not state.memory_scopes:
        sync_active_memory_scope(state)
    # Only expose the memory for the currently active scope.
    activate_memory_scope(state, preserve_current=False)
    return state



# ---------------------------------------------------------------------------
# Personal document library + local RAG
# ---------------------------------------------------------------------------


def default_document_library() -> dict[str, Any]:
    return {
        "projects": {
            "ROV_Control": {"folders": [str(ROOT / "documents" / "ROV_Control")], "recursive": True},
            "AEROSUB": {"folders": [str(ROOT / "documents" / "AEROSUB")], "recursive": True},
            "Thesis": {"folders": [str(ROOT / "documents" / "Thesis")], "recursive": True},
            "General": {"folders": [str(ROOT / "documents" / "General")], "recursive": True},
        }
    }

def ensure_document_library() -> dict[str, Any]:
    if not DOCUMENT_CONFIG_PATH.exists():
        DOCUMENT_CONFIG_PATH.write_text(yaml.safe_dump(default_document_library(), sort_keys=False), encoding="utf-8")
    try:
        data = yaml.safe_load(DOCUMENT_CONFIG_PATH.read_text(encoding="utf-8")) or {}
    except Exception as exc:
        print(f"[RAG] Could not read {DOCUMENT_CONFIG_PATH}: {exc}")
        data = default_document_library()
    if not isinstance(data.get("projects"), dict):
        data["projects"] = {}
    return data


def document_projects() -> dict[str, Any]:
    return ensure_document_library().get("projects", {})


def project_config(name: str) -> dict[str, Any] | None:
    cfg = document_projects().get(name)
    return cfg if isinstance(cfg, dict) else None


def init_rag_db() -> None:
    RAG_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(RAG_DB_PATH)
    try:
        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("""CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT, project TEXT NOT NULL, path TEXT NOT NULL,
            filename TEXT NOT NULL, ext TEXT, size_bytes INTEGER, mtime_ns INTEGER, sha256 TEXT,
            indexed_at TEXT, status TEXT DEFAULT 'indexed',
            page_count INTEGER DEFAULT 0, chunk_count INTEGER DEFAULT 0, extraction_method TEXT DEFAULT '',
            UNIQUE(project, path)
        )""")
        # Migrate existing databases created by earlier assistant versions.
        existing_cols={row[1] for row in conn.execute("PRAGMA table_info(documents)").fetchall()}
        for name, ddl in (
            ("page_count", "ALTER TABLE documents ADD COLUMN page_count INTEGER DEFAULT 0"),
            ("chunk_count", "ALTER TABLE documents ADD COLUMN chunk_count INTEGER DEFAULT 0"),
            ("extraction_method", "ALTER TABLE documents ADD COLUMN extraction_method TEXT DEFAULT ''"),
        ):
            if name not in existing_cols:
                conn.execute(ddl)
        conn.execute("""CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT, document_id INTEGER NOT NULL, project TEXT NOT NULL,
            page INTEGER, section TEXT, chunk_index INTEGER, text TEXT NOT NULL, sha256 TEXT,
            FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
        )""")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_chunks_project ON chunks(project)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_project ON documents(project)")
        try:
            conn.execute("CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(text, project, section, content='chunks', content_rowid='id')")
        except sqlite3.OperationalError:
            pass
        conn.execute("""CREATE TABLE IF NOT EXISTS chunk_embeddings (
            chunk_id INTEGER PRIMARY KEY, dim INTEGER NOT NULL, vector BLOB NOT NULL, model TEXT NOT NULL,
            FOREIGN KEY(chunk_id) REFERENCES chunks(id) ON DELETE CASCADE
        )""")
        conn.commit()
    finally:
        conn.close()


def _file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _extract_pdf_native(path: Path) -> tuple[list[tuple[int, str, str]], int]:
    """Extract text with pypdf and return (pages_with_text, total_pages)."""
    from pypdf import PdfReader
    out=[]
    reader=PdfReader(str(path))
    for i,page in enumerate(reader.pages,start=1):
        try:
            text=page.extract_text() or ""
        except Exception:
            text=""
        text=text.strip()
        if text:
            out.append((i,"",text))
    return out,len(reader.pages)


def _ocr_pdf_pages(path: Path, native_pages: dict[int,str], total_pages: int) -> tuple[list[tuple[int,str,str]], int, str]:
    """Use PyMuPDF + Tesseract only where native PDF text is missing/thin."""
    if not RAG_OCR_ENABLED:
        return [(p,"",t) for p,t in sorted(native_pages.items())],0,"native"
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
        import io
    except Exception as exc:
        print(f"[RAG][OCR] OCR dependencies unavailable: {exc}")
        return [(p,"",t) for p,t in sorted(native_pages.items())],0,"native-only"

    coverage=len(native_pages)/max(1,total_pages)
    avg_chars=(sum(len(x) for x in native_pages.values())/max(1,total_pages))
    needs_full_ocr = coverage < RAG_OCR_MIN_NATIVE_COVERAGE or avg_chars < RAG_OCR_MIN_PAGE_CHARS
    max_pages=RAG_OCR_MAX_PAGES if RAG_OCR_MAX_PAGES>0 else total_pages
    ocr_count=0
    out=[]
    doc=pymupdf.open(str(path))
    try:
        for page_no in range(1,total_pages+1):
            native=native_pages.get(page_no,"").strip()
            # OCR pages that are empty/thin; for mostly scanned PDFs OCR the whole document.
            should_ocr = needs_full_ocr or len(native)<RAG_OCR_MIN_PAGE_CHARS
            if not should_ocr or ocr_count>=max_pages:
                if native:
                    out.append((page_no,"",native))
                continue
            try:
                page=doc.load_page(page_no-1)
                scale=RAG_OCR_DPI/72.0
                pix=page.get_pixmap(matrix=pymupdf.Matrix(scale,scale),alpha=False)
                img=Image.open(io.BytesIO(pix.tobytes("png")))
                text=pytesseract.image_to_string(img,lang=RAG_OCR_LANG,timeout=RAG_OCR_TIMEOUT_SECONDS).strip()
                ocr_count+=1
                if text:
                    out.append((page_no,"",text))
                elif native:
                    out.append((page_no,"",native))
            except Exception as exc:
                if RAG_VERBOSE_LOGGING:
                    print(f"[RAG][OCR] page {page_no} failed: {exc}")
                if native:
                    out.append((page_no,"",native))
    finally:
        doc.close()
    method="ocr" if ocr_count else "native"
    if ocr_count and native_pages:
        method="hybrid-ocr"
    return out,ocr_count,method


def _extract_pdf(path: Path) -> tuple[list[tuple[int,str,str]], int, str]:
    native,total=_extract_pdf_native(path)
    native_map={page:text for page,_,text in native}
    sections,ocr_count,method=_ocr_pdf_pages(path,native_map,total)
    return sections,total,method


def _extract_docx(path: Path) -> list[tuple[int, str, str]]:
    from docx import Document
    doc = Document(str(path))
    paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = getattr(getattr(p, "style", None), "name", "") or ""
        paragraphs.append((0, style, text))
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        if rows:
            paragraphs.append((0, "table", "\n".join(rows)))
    return paragraphs


def _extract_plain(path: Path) -> list[tuple[int, str, str]]:
    return [(0, "", path.read_text(encoding="utf-8", errors="ignore"))]


def extract_document_sections(path: Path) -> list[tuple[int, str, str]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in {".txt", ".md"}:
        return _extract_plain(path)
    return []


def _split_text(text: str, size: int = RAG_CHUNK_SIZE, overlap: int = RAG_CHUNK_OVERLAP) -> list[str]:
    text = re.sub(r"\\s+", " ", text).strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    chunks=[]
    start=0
    while start < len(text):
        end=min(len(text), start+size)
        if end < len(text):
            cut=max(text.rfind(". ", start, end), text.rfind("; ", start, end), text.rfind(" ", start, end))
            if cut > start + int(size*0.55):
                end=cut+1
        chunk=text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start=max(0, end-overlap)
    return chunks


def _derive_section(text: str, fallback: str = "") -> str:
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    for line in lines[:8]:
        if 3 <= len(line) <= 140 and (line.endswith(":") or re.match(r"^(?:\\d+(?:\\.\\d+)*|Chapter|Appendix|Abstract|Introduction|Methods|Results|Discussion|Conclusion)\\b", line, re.I)):
            return line[:140]
    return fallback


def _index_file(conn: sqlite3.Connection, project: str, path: Path) -> dict[str, Any]:
    stat=path.stat()
    digest=_file_sha256(path)
    existing=conn.execute("SELECT id, sha256, size_bytes, mtime_ns, chunk_count FROM documents WHERE project=? AND path=?", (project,str(path))).fetchone()
    # Retry incomplete legacy/failed indexing even if the file itself has not changed.
    if existing and existing[1]==digest and existing[2]==stat.st_size and existing[3]==stat.st_mtime_ns and int(existing[4] or 0)>0:
        return {"action":"unchanged","pages":0,"chunks":int(existing[4] or 0),"method":"cached"}
    changed=bool(existing)
    if existing:
        doc_id=existing[0]
        conn.execute("DELETE FROM chunks WHERE document_id=?",(doc_id,))
        conn.execute("DELETE FROM chunk_embeddings WHERE chunk_id NOT IN (SELECT id FROM chunks)")
        conn.execute("UPDATE documents SET filename=?, ext=?, size_bytes=?, mtime_ns=?, sha256=?, indexed_at=?, status='indexing', page_count=0, chunk_count=0, extraction_method='' WHERE id=?",
                     (path.name,path.suffix.lower(),stat.st_size,stat.st_mtime_ns,digest,datetime.now(timezone.utc).isoformat(),doc_id))
    else:
        cur=conn.execute("INSERT INTO documents(project,path,filename,ext,size_bytes,mtime_ns,sha256,indexed_at,status,page_count,chunk_count,extraction_method) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                         (project,str(path),path.name,path.suffix.lower(),stat.st_size,stat.st_mtime_ns,digest,datetime.now(timezone.utc).isoformat(),"indexing",0,0,""))
        doc_id=cur.lastrowid
    try:
        extraction_method="native"
        if path.suffix.lower()==".pdf":
            sections,page_count,extraction_method=_extract_pdf(path)
        else:
            sections=extract_document_sections(path)
            page_count=len({p for p,_,_ in sections if p}) or 0
            if path.suffix.lower() not in {".pdf"} and page_count==0:
                page_count=1 if sections and any(t.strip() for _,_,t in sections) else 0
        idx=0
        for page,section,text in sections:
            section_name=section or _derive_section(text)
            for chunk in _split_text(text):
                idx+=1
                chash=hashlib.sha256(chunk.encode("utf-8",errors="ignore")).hexdigest()
                conn.execute("INSERT INTO chunks(document_id,project,page,section,chunk_index,text,sha256) VALUES(?,?,?,?,?,?,?)",
                             (doc_id,project,page,section_name,idx,chunk,chash))
        status="indexed" if idx else "no extractable text"
        conn.execute("UPDATE documents SET indexed_at=?, status=?, page_count=?, chunk_count=?, extraction_method=? WHERE id=?",
                     (datetime.now(timezone.utc).isoformat(),status,page_count,idx,extraction_method,doc_id))
        return {"action":"changed" if changed else "new","pages":page_count,"chunks":idx,"method":extraction_method}
    except Exception as exc:
        conn.execute("UPDATE documents SET indexed_at=?, status=?, page_count=?, chunk_count=?, extraction_method=? WHERE id=?",
                     (datetime.now(timezone.utc).isoformat(),f"error: {exc}",0,0,"error",doc_id))
        return {"action":"error","pages":0,"chunks":0,"method":"error","error":str(exc)}


def scan_project_documents(project: str) -> dict[str, int]:
    init_rag_db()
    cfg=project_config(project)
    if not cfg:
        raise ValueError(f"Unknown document project: {project}")
    folders=cfg.get("folders",[]) if isinstance(cfg.get("folders",[]),list) else []
    recursive=bool(cfg.get("recursive",True))
    conn=sqlite3.connect(RAG_DB_PATH)
    stats={"new":0,"changed":0,"skipped":0,"errors":0,"seen":0,"indexed_files":0,"ocr_files":0,"chunks":0,"pages":0}
    seen=set()
    valid_root_found=False
    try:
        for raw in folders:
            folder=Path(os.path.expanduser(str(raw)))
            folder=(ROOT/folder).resolve() if not folder.is_absolute() else folder.resolve()
            if not folder.exists():
                print(f"[RAG] Missing folder for {project}: {folder}")
                continue
            valid_root_found=True
            iterator=folder.rglob("*") if recursive else folder.glob("*")
            for path in iterator:
                if not path.is_file() or path.suffix.lower() not in RAG_SUPPORTED_EXTENSIONS:
                    continue
                seen.add(str(path)); stats["seen"]+=1
                try:
                    if path.stat().st_size>RAG_MAX_FILE_MB*1024*1024:
                        stats["skipped"]+=1
                        print(f"[RAG] Skipping oversized file: {path.name} ({path.stat().st_size/1024/1024:.1f} MB)")
                        continue
                    result=_index_file(conn,project,path)
                    action=result.get("action")
                    if action=="new": stats["new"]+=1
                    elif action=="changed": stats["changed"]+=1
                    elif action=="error":
                        stats["errors"]+=1
                        print(f"[RAG] Failed to index {path.name}: {result.get('error','unknown error')}")
                    if int(result.get("chunks",0))>0:
                        stats["indexed_files"]+=1
                    stats["chunks"]+=int(result.get("chunks",0) or 0)
                    stats["pages"]+=int(result.get("pages",0) or 0)
                    if "ocr" in str(result.get("method","")):
                        stats["ocr_files"]+=1
                except Exception as exc:
                    stats["errors"]+=1
                    print(f"[RAG] Indexing error for {path.name}: {exc}")
        if valid_root_found:
            rows=conn.execute("SELECT id,path FROM documents WHERE project=?",(project,)).fetchall()
            for doc_id,path in rows:
                if path not in seen:
                    conn.execute("DELETE FROM chunks WHERE document_id=?",(doc_id,))
                    conn.execute("DELETE FROM documents WHERE id=?",(doc_id,))
        try:
            conn.execute("INSERT INTO chunks_fts(chunks_fts) VALUES('rebuild')")
        except sqlite3.OperationalError:
            pass
        conn.commit()
    finally:
        conn.close()
    return stats



def list_indexed_documents(project: str | None = None) -> list[dict[str, Any]]:
    init_rag_db()
    conn=sqlite3.connect(RAG_DB_PATH)
    try:
        if project:
            rows=conn.execute("SELECT project,filename,path,size_bytes,indexed_at,status,page_count,chunk_count,extraction_method FROM documents WHERE project=? ORDER BY filename", (project,)).fetchall()
        else:
            rows=conn.execute("SELECT project,filename,path,size_bytes,indexed_at,status,page_count,chunk_count,extraction_method FROM documents ORDER BY project,filename").fetchall()
        return [{"project":r[0],"filename":r[1],"path":r[2],"size_bytes":r[3],"indexed_at":r[4],"status":r[5],"page_count":r[6],"chunk_count":r[7],"extraction_method":r[8]} for r in rows]
    finally:
        conn.close()


def _ollama_embed(texts: list[str]) -> list[list[float]]:
    if not RAG_EMBEDDINGS_ENABLED or not texts:
        return []
    try:
        r=requests.post(f"{OLLAMA_BASE_URL}/api/embed", json={"model":RAG_EMBEDDING_MODEL,"input":texts}, timeout=180)
        r.raise_for_status()
        data=r.json()
        embeddings=data.get("embeddings")
        if isinstance(embeddings,list) and embeddings and isinstance(embeddings[0],list):
            return embeddings
    except Exception as exc:
        print(f"[RAG] Embedding model unavailable; using lexical retrieval: {exc}")
        return []
    return []


def _cosine(a: list[float], b: list[float]) -> float:
    if not a or not b or len(a)!=len(b): return 0.0
    dot=sum(x*y for x,y in zip(a,b)); na=math.sqrt(sum(x*x for x in a)); nb=math.sqrt(sum(y*y for y in b))
    return dot/(na*nb) if na and nb else 0.0


def rag_search(project: str, question: str, top_k: int = RAG_TOP_K) -> list[Result]:
    init_rag_db()
    scan_project_documents(project)
    conn=sqlite3.connect(RAG_DB_PATH)
    try:
        lexical=[]
        qtokens=sorted(tokens(question))
        if qtokens:
            query=" OR ".join(qtokens)
            try:
                rows=conn.execute("""SELECT c.id,c.project,c.page,c.section,c.text,d.filename,d.path,bm25(chunks_fts)
                    FROM chunks_fts JOIN chunks c ON chunks_fts.rowid=c.id JOIN documents d ON d.id=c.document_id
                    WHERE chunks_fts MATCH ? AND c.project=? ORDER BY bm25(chunks_fts) LIMIT ?""", (query,project,RAG_CANDIDATE_K)).fetchall()
                lexical=[r for r in rows]
            except sqlite3.OperationalError:
                rows=conn.execute("SELECT c.id,c.project,c.page,c.section,c.text,d.filename,d.path FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.project=? LIMIT ?", (project,RAG_CANDIDATE_K*2)).fetchall()
                lexical=sorted(rows,key=lambda r:score_text(question,r[4]),reverse=True)[:RAG_CANDIDATE_K]
        else:
            lexical=conn.execute("SELECT c.id,c.project,c.page,c.section,c.text,d.filename,d.path FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.project=? ORDER BY c.id DESC LIMIT ?",(project,RAG_CANDIDATE_K)).fetchall()
        qembs=_ollama_embed([question])
        vector_scores={}
        if qembs:
            qvec=qembs[0]
            ids=[int(r[0]) for r in lexical]
            if ids:
                marks=','.join('?'*len(ids))
                erows=conn.execute(f"SELECT chunk_id,dim,vector FROM chunk_embeddings WHERE chunk_id IN ({marks})",ids).fetchall()
                for cid,dim,blob in erows:
                    vector_scores[int(cid)]=_cosine(qvec, list(__import__('numpy').frombuffer(blob,dtype='float32')))
        scored=[]
        for i,r in enumerate(lexical):
            cid,proj,page,section,text,filename,path,*rest=r
            lexical_score=score_text(question,text)
            bm=0.0
            if rest:
                try: bm=1.0/(1.0+float(rest[0]))
                except: bm=0.0
            vec=vector_scores.get(int(cid),0.0)
            score=0.55*lexical_score+0.25*bm+0.20*vec if qembs else 0.75*lexical_score+0.25*bm
            scored.append((score,r))
        scored.sort(key=lambda x:x[0],reverse=True)
        out=[]
        for score,r in scored[:top_k]:
            cid,proj,page,section,text,filename,path,*_=r
            out.append(Result("document",f"{filename} — {section or 'document'}",path,text[:4000],str(page or ""),{"project":project,"page":page,"section":section,"score":round(score,4),"chunk_id":cid}))
        return out
    finally:
        conn.close()


def print_document_projects(current: str | None = None) -> None:
    projects=document_projects()
    print("\nDocument projects")
    print("="*80)
    if not projects:
        print("No document projects configured.")
    for idx,(name,cfg) in enumerate(projects.items(),1):
        rows=list_indexed_documents(name)
        count=sum(1 for row in rows if int(row.get("chunk_count") or 0)>0)
        chunks=sum(int(row.get("chunk_count") or 0) for row in rows)
        mark="*" if name==current else " "
        folders=cfg.get("folders",[]) if isinstance(cfg,dict) else []
        print(f"{mark} {idx}. {name:22} {count:>5} indexed | {chunks:>6} chunks | {len(folders)} folder(s)")
    print("="*80)


def choose_document_project(current: str | None = None) -> str | None:
    projects=document_projects()
    names=list(projects)
    if not names:
        return None
    while True:
        print_document_projects(current)
        print("0. None / disable document RAG")
        choice=input("Select project number: ").strip()
        if choice=="0": return None
        try:
            idx=int(choice)-1
            if 0<=idx<len(names): return names[idx]
        except ValueError: pass
        print("Invalid selection.")


def document_manager(current: str | None = None) -> str | None:
    while True:
        print("\nDocument manager")
        print("  1. Scan all projects")
        print("  2. Scan current project")
        print("  3. Show indexed documents")
        print("  4. Change document project")
        print("  5. Add project/folder to YAML")
        print("  q. Back")
        choice=input("Select: ").strip().lower()
        if choice=="q": return current
        if choice=="1":
            for name in document_projects():
                print(f"[RAG] {name}: {scan_project_documents(name)}")
        elif choice=="2":
            if current: print(f"[RAG] {current}: {scan_project_documents(current)}")
            else: print("No current document project selected.")
        elif choice=="3":
            rows=list_indexed_documents(current)
            print("\nIndexed documents")
            for row in rows:
                print(f"[{row['project']}] {row['filename']} | pages={row.get('page_count',0)} | chunks={row.get('chunk_count',0)} | {row.get('extraction_method','')} | {row['status']}")
            print(f"Total: {len(rows)}")
        elif choice=="4":
            current=choose_document_project(current)
        elif choice=="5":
            name=input("Project name: ").strip()
            folder=input("Folder path: ").strip()
            if name and folder:
                data=ensure_document_library(); data.setdefault("projects",{}).setdefault(name,{"folders":[],"recursive":True})
                if folder not in data["projects"][name]["folders"]: data["projects"][name]["folders"].append(folder)
                DOCUMENT_CONFIG_PATH.write_text(yaml.safe_dump(data,sort_keys=False),encoding="utf-8")
                print(f"Added folder to {name} in {DOCUMENT_CONFIG_PATH}")
        else: print("Invalid selection.")
# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------


def parse_run_stamp(path: Path) -> datetime:
    try:
        return datetime.strptime(path.name, "%Y%m%d_%H%M%S").replace(tzinfo=timezone.utc)
    except Exception:
        return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)


def infer_mode(summary: dict[str, Any]) -> str:
    mode = str(summary.get("mode", "")).lower()
    if mode in {"pilot", "expanded", "full", "quick"}:
        return "expanded" if mode == "full" else mode
    minimum = int(summary.get("minimum_candidate_papers", 150) or 150)
    deep_read = int(summary.get("deep_read_per_task", 50) or 50)
    if minimum <= 20 or deep_read <= 5:
        return "pilot"
    return "expanded"


def load_run_record(run_dir: Path) -> Optional[dict[str, Any]]:
    summary_path = run_dir / "run_summary.json"
    if not summary_path.exists():
        return None
    try:
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    except Exception:
        return None
    stamp = parse_run_stamp(run_dir)
    now = datetime.now(timezone.utc)
    age_days = max(0.0, (now - stamp).total_seconds() / 86400.0)
    mode = infer_mode(summary)
    status = summary.get("status") or ("completed" if summary.get("pdf_verification", {}).get("ok") else "completed-no-verified-pdf")
    if mode != "quick" and not summary.get("paper_counts"):
        status = "incomplete"
    if mode == "pilot":
        refresh = "needs expanded run"
    elif mode == "quick":
        refresh = "quick scan — expand to deep research"
    elif age_days <= RECENT_DAYS:
        refresh = "recent"
    elif age_days <= 30:
        refresh = "consider refresh"
    elif age_days <= 90:
        refresh = "aging — refresh useful"
    else:
        refresh = "stale — rerun recommended"

    return {
        "id": run_dir.name,
        "title": summary.get("title") or summary.get("question") or run_dir.name,
        "question": summary.get("question", ""),
        "created": stamp.isoformat(),
        "age_days": round(age_days, 1),
        "mode": mode,
        "status": status,
        "refresh": refresh,
        "run_dir": str(run_dir),
        "pdf": summary.get("pdf", ""),
        "paper_counts": summary.get("paper_counts", {}),
        "source_counts": summary.get("source_counts", {}),
        "deep_read_per_task": summary.get("deep_read_per_task"),
    }


def build_registry() -> list[dict[str, Any]]:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    records = []
    for run_dir in RUNS_DIR.iterdir():
        if not run_dir.is_dir():
            continue
        record = load_run_record(run_dir)
        if record:
            records.append(record)
    records.sort(key=lambda x: x["created"], reverse=True)
    REGISTRY_PATH.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    return records


def print_registry_page(records: list[dict[str, Any]], offset: int = 0) -> None:
    page = records[offset:offset + PAGE_SIZE]
    if not page:
        print("No more research topics.")
        return
    print("\nRecent research topics")
    print("=" * 110)
    for i, r in enumerate(page, start=1 + offset):
        print(
            f"[{i:>2}] {r['title'][:58]:58} | {r['created'][:10]} | "
            f"{r['age_days']:>6.1f}d | {r['mode']:<8} | {r['status']:<25} | {r['refresh']}"
        )
    print("=" * 110)


def choose_run(records: list[dict[str, Any]], requested: str | None = None) -> dict[str, Any] | None:
    """Interactive selector: latest 10 projects + item 11 for a new topic."""
    if requested:
        q = requested.strip().lower()
        for r in records:
            if r["id"].lower() == q:
                return r
        matches = [r for r in records if q in r["title"].lower() or q in r["question"].lower()]
        if len(matches) == 1:
            return matches[0]
        if matches:
            records = matches
        else:
            raise ValueError(f"No research project matched: {requested}")

    offset = 0
    while True:
        page = records[offset:offset + PAGE_SIZE]
        print("\nRecent research topics")
        print("=" * 118)
        if page:
            for i, r in enumerate(page, start=1 + offset):
                print(
                    f"[{i:>2}] {r['title'][:54]:54} | {r['created'][:10]} | "
                    f"{r['age_days']:>6.1f}d | {r['mode']:<8} | {r['status']:<18} | {r['refresh']}"
                )
        else:
            print("No saved research topics yet.")
        print("-" * 118)
        print("[11] Open New Topic")
        print("      Search the internet + local LLM without a deep-research corpus.")
        print("=" * 118)

        max_idx = min(offset + PAGE_SIZE, len(records))
        extras = []
        if offset + PAGE_SIZE < len(records):
            extras.append("m=more")
        prompt = f"Select [{offset + 1}-{max_idx}]" if page else "Select"
        prompt += ", 11=new topic"
        if extras:
            prompt += ", m=more"
        prompt += ", q=quit: "

        try:
            choice = input(prompt).strip().lower()
        except EOFError:
            raise SystemExit(0)

        if choice == "q":
            raise SystemExit(0)
        if choice == "11":
            return None
        if choice == "m":
            if offset + PAGE_SIZE >= len(records):
                print("No older topics remain.")
            else:
                offset += PAGE_SIZE
            continue
        try:
            idx = int(choice)
        except ValueError:
            print("Please enter a topic number, 11, m, or q.")
            continue
        if 1 <= idx <= len(records):
            return records[idx - 1]
        print("Selection out of range.")


# ---------------------------------------------------------------------------
# Local corpus search
# ---------------------------------------------------------------------------

TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9+_.-]{2,}")
STOP = {
    "what", "which", "where", "when", "does", "have", "with", "from", "that", "this",
    "those", "into", "about", "there", "their", "should", "could", "would", "than", "been",
    "were", "using", "used", "under", "over", "for", "and", "the", "are", "was", "can", "how",
}


def tokens(text: str) -> set[str]:
    return {x.lower() for x in TOKEN_RE.findall(text) if x.lower() not in STOP}


def score_text(query: str, text: str) -> float:
    q = tokens(query)
    t = tokens(text)
    if not q or not t:
        return 0.0
    overlap = len(q & t) / len(q)
    phrase = fuzz_ratio(query.lower(), text[:1000].lower()) / 100.0
    return 0.7 * overlap + 0.3 * phrase


def local_results(run: dict[str, Any], question: str) -> dict[str, list[Result]]:
    out = {"papers": [], "evidence": [], "technical": []}
    run_dir = Path(run["run_dir"])
    db_path = run_dir / "papers.sqlite"
    if db_path.exists():
        conn = sqlite3.connect(db_path)
        try:
            rows = conn.execute(
                "SELECT title, abstract, year, authors, venue, doi, url, source, task FROM papers"
            ).fetchall()
            scored = []
            for title, abstract, year, authors, venue, doi, url, source, task in rows:
                blob = f"{title} {abstract or ''} {authors or ''} {venue or ''} {task or ''}"
                scored.append((score_text(question, blob), title, abstract or "", url or "", year, source, task, doi or ""))
            scored.sort(reverse=True, key=lambda x: x[0])
            for s, title, abstract, url, year, source, task, doi in scored[:LOCAL_PAPER_RESULTS]:
                out["papers"].append(Result(
                    source=f"local:{source}", title=title, url=url,
                    snippet=abstract[:900], date=str(year or ""),
                    metadata={"score": round(s, 4), "task": task, "doi": doi},
                ))

            tech_rows = conn.execute(
                "SELECT title, url, source, snippet, query, pass_name FROM technical_sources"
            ).fetchall()
            tech_scored = []
            for title, url, source, snippet, query, pass_name in tech_rows:
                blob = f"{title} {snippet or ''} {query or ''}"
                tech_scored.append((score_text(question, blob), title, url, source, snippet or "", query or "", pass_name or ""))
            tech_scored.sort(reverse=True, key=lambda x: x[0])
            for s, title, url, source, snippet, query, pass_name in tech_scored[:LOCAL_TECH_RESULTS]:
                out["technical"].append(Result(
                    source=f"local:{source}", title=title, url=url, snippet=snippet,
                    metadata={"score": round(s, 4), "query": query, "pass_name": pass_name},
                ))
        finally:
            conn.close()

    for name in ("evidence_cards_post_gap.json", "evidence_cards.json"):
        path = run_dir / name
        if not path.exists():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        cards = []
        for task, values in data.items():
            for card in values or []:
                if not isinstance(card, dict):
                    continue
                title = str(card.get("title") or card.get("paper_title") or "")
                blob = json.dumps(card, ensure_ascii=False)
                cards.append((score_text(question, blob), title, blob[:2500], card.get("url", ""), task))
        cards.sort(reverse=True, key=lambda x: x[0])
        for s, title, snippet, url, task in cards[:LOCAL_EVIDENCE_RESULTS]:
            out["evidence"].append(Result(
                source="local:evidence",
                title=title or task,
                url=url,
                snippet=snippet,
                metadata={"score": round(s, 4), "task": task, "artifact": name},
            ))
        break
    return out


# ---------------------------------------------------------------------------
# External sources
# ---------------------------------------------------------------------------


def disable_source(name: str, reason: str) -> None:
    SOURCE_DISABLED[name] = True
    SOURCE_FAILURE_REASONS[name] = str(reason)[:300]
    terminal_log(name.upper(), f"disabled for this question: {reason}", "yellow")


def source_ready(name: str) -> bool:
    cfg = SOURCE_CONFIG[name]
    if not cfg["enabled"] or SOURCE_DISABLED[name]:
        return False
    if SOURCE_ERRORS[name] >= cfg["failures"]:
        SOURCE_DISABLED[name] = True
        return False
    return True


def limited_get(name: str, url: str, *, params: dict[str, Any] | None = None,
                 headers: dict[str, str] | None = None, timeout: float = 30.0) -> requests.Response:
    cfg = SOURCE_CONFIG[name]
    with SOURCE_LOCKS[name]:
        if not source_ready(name):
            raise RuntimeError(f"source disabled: {name}")
        wait = cfg["interval"] - (time.monotonic() - SOURCE_LAST[name])
        if wait > 0:
            time.sleep(wait)
        response = requests.get(url, params=params, headers=headers or HEADERS, timeout=timeout)
        SOURCE_LAST[name] = time.monotonic()
        if response.status_code in {408, 425, 428, 429, 500, 502, 503, 504}:
            SOURCE_ERRORS[name] += 1
            raise requests.HTTPError(f"HTTP {response.status_code}", response=response)
        if response.status_code in {401, 403}:
            SOURCE_ERRORS[name] += 1
            raise requests.HTTPError(f"HTTP {response.status_code}", response=response)
        response.raise_for_status()
        SOURCE_ERRORS[name] = 0
        return response


def web_search(query: str, limit: int) -> list[Result]:
    if DDGS is None:
        return []
    with SOURCE_LOCKS["web"]:
        if not source_ready("web"):
            return []
        wait = SOURCE_CONFIG["web"]["interval"] - (time.monotonic() - SOURCE_LAST["web"])
        if wait > 0:
            time.sleep(wait)
        try:
            data = DDGS().text(query, max_results=limit)
            SOURCE_LAST["web"] = time.monotonic()
            SOURCE_ERRORS["web"] = 0
            return [Result("web", x.get("title", ""), x.get("href", ""), x.get("body", "")) for x in data]
        except Exception as exc:
            SOURCE_LAST["web"] = time.monotonic()
            SOURCE_ERRORS["web"] += 1
            SOURCE_FAILURE_REASONS["web"] = str(exc)[:300]
            if SOURCE_ERRORS["web"] >= SOURCE_CONFIG["web"]["failures"]:
                disable_source("web", str(exc))
            return []


def google_search(query: str, limit: int) -> list[Result]:
    if SERPAPI_API_KEY:
        try:
            r = limited_get(
                "google",
                "https://serpapi.com/search.json",
                params={"engine": "google", "q": query, "api_key": SERPAPI_API_KEY, "num": min(limit, 10)},
            )
            data = r.json()
            return [
                Result("google", x.get("title", ""), x.get("link", ""), x.get("snippet", ""),
                       metadata={"position": x.get("position")})
                for x in data.get("organic_results", [])
            ]
        except Exception as exc:
            SOURCE_ERRORS["google"] += 1
            SOURCE_FAILURE_REASONS["google"] = str(exc)[:300]
            if SOURCE_ERRORS["google"] >= SOURCE_CONFIG["google"]["failures"]:
                disable_source("google", str(exc))
            return []
    return [Result(r.source + ":fallback", r.title, r.url, r.snippet, r.date, r.metadata)
            for r in web_search(query, limit)]


def youtube_search(query: str, limit: int) -> list[Result]:
    if not YOUTUBE_API_KEY:
        return [Result("youtube:web", r.title, r.url, r.snippet, r.date, r.metadata)
                for r in web_search(f"site:youtube.com {query}", limit)]
    try:
        r = limited_get(
            "youtube",
            "https://www.googleapis.com/youtube/v3/search",
            params={"part": "snippet", "q": query, "type": "video", "maxResults": min(limit, 10), "key": YOUTUBE_API_KEY},
        )
        data = r.json()
        out = []
        for x in data.get("items", []):
            vid = ((x.get("id") or {}).get("videoId") or "")
            sn = x.get("snippet") or {}
            if vid:
                out.append(Result("youtube", sn.get("title", ""), f"https://www.youtube.com/watch?v={vid}",
                                  sn.get("description", ""), sn.get("publishedAt", ""),
                                  {"channel": sn.get("channelTitle", "") }))
        return out
    except Exception as exc:
        SOURCE_ERRORS["youtube"] += 1
        if SOURCE_ERRORS["youtube"] >= SOURCE_CONFIG["youtube"]["failures"]:
            disable_source("youtube", str(exc))
        return []


def github_search(query: str, limit: int) -> list[Result]:
    headers = dict(HEADERS)
    headers["Accept"] = "application/vnd.github+json"
    headers["X-GitHub-Api-Version"] = "2026-03-10"
    if GITHUB_TOKEN:
        headers["Authorization"] = f"Bearer {GITHUB_TOKEN}"
    try:
        r = limited_get(
            "github",
            "https://api.github.com/search/repositories",
            params={"q": query, "per_page": min(limit, 10), "sort": "stars", "order": "desc"},
            headers=headers,
        )
        data = r.json()
        return [Result(
            "github", x.get("full_name", ""), x.get("html_url", ""), x.get("description", ""),
            x.get("updated_at", ""),
            {"stars": x.get("stargazers_count", 0), "forks": x.get("forks_count", 0), "language": x.get("language", "")},
        ) for x in data.get("items", [])]
    except Exception as exc:
        SOURCE_ERRORS["github"] += 1
        if SOURCE_ERRORS["github"] >= SOURCE_CONFIG["github"]["failures"]:
            disable_source("github", str(exc))
        return []


def reddit_search(query: str, limit: int) -> list[Result]:
    return [Result("reddit:web", r.title, r.url, r.snippet, r.date, r.metadata)
            for r in web_search(f"site:reddit.com {query}", limit)]


def quora_search(query: str, limit: int) -> list[Result]:
    return [Result("quora:web", r.title, r.url, r.snippet, r.date, r.metadata)
            for r in web_search(f"site:quora.com {query}", limit)]


def openalex_search(query: str, limit: int) -> list[Result]:
    try:
        params = {"search": query, "per-page": min(limit, 25),
                  "select": "id,title,publication_year,doi,primary_location,cited_by_count"}
        if OPENALEX_API_KEY:
            params["api_key"] = OPENALEX_API_KEY
        r = limited_get("openalex", "https://api.openalex.org/works", params=params)
        data = r.json()
        out = []
        for x in data.get("results", []):
            loc = x.get("primary_location") or {}
            out.append(Result(
                "openalex", x.get("title", ""), x.get("doi") or x.get("id") or "",
                f"Citations: {x.get('cited_by_count', 0)}",
                str(x.get("publication_year") or ""),
                {"doi": x.get("doi", ""), "cited_by_count": x.get("cited_by_count", 0),
                 "source": ((loc.get("source") or {}).get("display_name") if isinstance(loc.get("source"), dict) else "")},
            ))
        return out
    except Exception as exc:
        SOURCE_ERRORS["openalex"] += 1
        if SOURCE_ERRORS["openalex"] >= SOURCE_CONFIG["openalex"]["failures"]:
            disable_source("openalex", str(exc))
        return []


def semantic_scholar_search(query: str, limit: int) -> list[Result]:
    if not SEMANTIC_SCHOLAR_ENABLED:
        return []
    headers = dict(HEADERS)
    if SEMANTIC_SCHOLAR_API_KEY:
        headers["x-api-key"] = SEMANTIC_SCHOLAR_API_KEY
    try:
        r = limited_get(
            "semantic_scholar",
            "https://api.semanticscholar.org/graph/v1/paper/search",
            params={"query": query, "limit": min(limit, 10), "fields": "title,year,authors,url,externalIds,citationCount,abstract"},
            headers=headers,
        )
        data = r.json()
        out = []
        for x in data.get("data", []):
            ids = x.get("externalIds") or {}
            out.append(Result(
                "semantic_scholar", x.get("title", ""), x.get("url", ""), x.get("abstract", "")[:900],
                str(x.get("year") or ""), {"citations": x.get("citationCount", 0), "doi": ids.get("DOI", "")},
            ))
        return out
    except Exception as exc:
        SOURCE_ERRORS["semantic_scholar"] += 1
        if SOURCE_ERRORS["semantic_scholar"] >= SOURCE_CONFIG["semantic_scholar"]["failures"]:
            disable_source("semantic_scholar", str(exc))
        return []


def arxiv_search(query: str, limit: int) -> list[Result]:
    try:
        r = limited_get(
            "arxiv",
            "https://export.arxiv.org/api/query",
            params={"search_query": f"all:{query}", "start": 0, "max_results": min(limit, 10), "sortBy": "relevance"},
        )
        root = ET.fromstring(r.text)
        ns = {"a": "http://www.w3.org/2005/Atom"}
        out = []
        for entry in root.findall("a:entry", ns):
            title = " ".join((entry.findtext("a:title", default="", namespaces=ns) or "").split())
            summary = " ".join((entry.findtext("a:summary", default="", namespaces=ns) or "").split())
            url = entry.findtext("a:id", default="", namespaces=ns)
            published = entry.findtext("a:published", default="", namespaces=ns)
            out.append(Result("arxiv", title, url, summary[:1000], published, {}))
        return out
    except Exception as exc:
        SOURCE_ERRORS["arxiv"] += 1
        if SOURCE_ERRORS["arxiv"] >= SOURCE_CONFIG["arxiv"]["failures"]:
            disable_source("arxiv", str(exc))
        return []


def crossref_search(query: str, limit: int) -> list[Result]:
    try:
        params = {"query.bibliographic": query, "rows": min(limit, 10),
                  "select": "DOI,title,author,published,URL,container-title"}
        r = limited_get("crossref", "https://api.crossref.org/works", params=params)
        data = r.json()
        out = []
        for x in (data.get("message") or {}).get("items", []):
            titles = x.get("title") or []
            out.append(Result(
                "crossref", titles[0] if titles else "", x.get("URL", ""),
                x.get("container-title", [""])[0] if x.get("container-title") else "",
                str(((x.get("published-print") or x.get("published-online") or {}).get("date-parts") or [[""]])[0][0]),
                {"doi": x.get("DOI", "")},
            ))
        return out
    except Exception as exc:
        SOURCE_ERRORS["crossref"] += 1
        if SOURCE_ERRORS["crossref"] >= SOURCE_CONFIG["crossref"]["failures"]:
            disable_source("crossref", str(exc))
        return []


def dblp_search(query: str, limit: int) -> list[Result]:
    try:
        # DBLP Search API returns JSON with this format endpoint.
        url = "https://dblp.org/search/publ/api"
        r = limited_get("dblp", url, params={"q": query, "h": min(limit, 10), "format": "json"})
        data = r.json()
        hits = (((data.get("result") or {}).get("hits") or {}).get("hit") or [])
        if isinstance(hits, dict):
            hits = [hits]
        out = []
        for h in hits:
            info = h.get("info") or {}
            out.append(Result(
                "dblp", info.get("title", ""), info.get("ee") or info.get("url") or "",
                info.get("venue", ""), str(info.get("year", "")),
                {"authors": info.get("authors", {}).get("author", []) if isinstance(info.get("authors"), dict) else info.get("authors", "")},
            ))
        return out
    except Exception as exc:
        SOURCE_ERRORS["dblp"] += 1
        if SOURCE_ERRORS["dblp"] >= SOURCE_CONFIG["dblp"]["failures"]:
            disable_source("dblp", str(exc))
        return []


# ---------------------------------------------------------------------------
# Local LLM helpers
# ---------------------------------------------------------------------------


def invoke_model(model_name: str, prompt: str, *, ctx: int, tokens_out: int, heartbeat: float) -> str:
    started=time.monotonic()
    terminal_log("LLM", f"{model_name} | waiting...", "magenta")
    if ChatOllama is not None:
        try:
            llm=ChatOllama(model=model_name,temperature=0,num_ctx=ctx,num_predict=tokens_out,
                           client_kwargs={"timeout":LLM_TIMEOUT},reasoning=False)
            result=llm.invoke(prompt)
            content=result.content if hasattr(result,"content") else str(result)
            terminal_log("LLM", f"{model_name} | completed in {time.monotonic()-started:.1f}s", "green")
            return content
        except Exception as exc:
            print(f"[LLM] LangChain Ollama failed; trying direct Ollama HTTP: {exc}")
    payload={"model":model_name,"messages":[{"role":"user","content":prompt}],"stream":False,
             "options":{"temperature":0,"num_ctx":ctx,"num_predict":tokens_out}}
    try:
        r=requests.post(f"{OLLAMA_BASE_URL}/api/chat",json=payload,timeout=LLM_TIMEOUT)
        r.raise_for_status()
        data=r.json()
        content=((data.get("message") or {}).get("content")) or data.get("response") or ""
        terminal_log("LLM", f"{model_name} | completed in {time.monotonic()-started:.1f}s", "green")
        return str(content)
    except Exception as exc:
        raise RuntimeError(f"Ollama model call failed for {model_name}: {exc}") from exc


def make_query_plan(question: str, run: dict[str, Any], conversation_context: str = "") -> dict[str, Any]:
    prompt = f"""
You are the search strategist for a follow-up research question.

Research project:
{run['title']}
Original question:
{run['question']}
Project data mode: {run['mode']}
Project age: {run['age_days']} days

Follow-up question:
{question}

Conversation memory for continuity (not evidence):
{conversation_context[:20000]}

Create compact, high-signal search queries for these source families:
1. general web / Google
2. academic literature
3. YouTube
4. GitHub
5. Reddit
6. Quora
7. adjacent technical sources

Return JSON with keys:
web_queries, academic_queries, youtube_queries, github_queries,
reddit_queries, quora_queries, technical_queries.
Use 1-3 distinct queries per family. Do not write explanations.
"""
    text = invoke_model(QUERY_MODEL, prompt, ctx=QUERY_CTX, tokens_out=QUERY_TOKENS, heartbeat=ANSWER_HEARTBEAT)
    try:
        m = re.search(r"\{.*\}", text, flags=re.S)
        if not m:
            raise ValueError("No JSON object found")
        data = json.loads(m.group(0))
        return {k: list(dict.fromkeys(data.get(k, [])))[:FOLLOWUP_QUERIES_PER_FAMILY] for k in (
            "web_queries", "academic_queries", "youtube_queries", "github_queries",
            "reddit_queries", "quora_queries", "technical_queries")}
    except Exception:
        return {
            "web_queries": [question],
            "academic_queries": [question],
            "youtube_queries": [question],
            "github_queries": [question],
            "reddit_queries": [question],
            "quora_queries": [question],
            "technical_queries": [question],
        }


def run_external_searches(plan: dict[str, Any]) -> list[Result]:
    global LAST_EXTERNAL_STATUS
    started = time.monotonic()
    jobs = []
    for q in plan.get("web_queries", [])[:FOLLOWUP_QUERIES_PER_FAMILY]:
        jobs.append(("google", google_search, q))
    for q in plan.get("academic_queries", [])[:FOLLOWUP_QUERIES_PER_FAMILY]:
        jobs.extend([
            ("openalex", openalex_search, q),
            ("semantic_scholar", semantic_scholar_search, q),
            ("arxiv", arxiv_search, q),
            ("crossref", crossref_search, q),
            ("dblp", dblp_search, q),
        ])
    for q in plan.get("youtube_queries", [])[:1]:
        jobs.append(("youtube", youtube_search, q))
    for q in plan.get("github_queries", [])[:1]:
        jobs.append(("github", github_search, q))
    for q in plan.get("reddit_queries", [])[:1]:
        jobs.append(("reddit", reddit_search, q))
    for q in plan.get("quora_queries", [])[:1]:
        jobs.append(("quora", quora_search, q))
    for q in plan.get("technical_queries", [])[:1]:
        jobs.append(("web-technical", google_search, q))

    def execute(job):
        label, fn, q = job
        t0 = time.monotonic()
        try:
            results = fn(q, EXTERNAL_RESULTS_PER_SOURCE)
            for r in results:
                r.metadata = dict(r.metadata or {})
                r.metadata["query"] = q
                r.metadata["source_family"] = label
            return label, results, time.monotonic() - t0, ""
        except Exception as exc:
            return label, [], time.monotonic() - t0, str(exc)

    results: list[Result] = []
    source_counts: dict[str, int] = {}
    worker_failures: dict[str, str] = {}
    if jobs:
        terminal_log("WEB", f"starting {len(jobs)} external search job(s)...", "blue")
    with concurrent.futures.ThreadPoolExecutor(max_workers=EXTERNAL_MAX_WORKERS) as pool:
        futures = [pool.submit(execute, job) for job in jobs]
        for fut in concurrent.futures.as_completed(futures):
            label, values, elapsed, error = fut.result()
            source_counts[label] = source_counts.get(label, 0) + len(values)
            results.extend(values)
            if error:
                worker_failures[label] = error[:300]
            terminal_log("WEB", f"{label} -> {len(values)} result(s) in {elapsed:.1f}s", "dim")

    unique = deduplicate_results(results)
    elapsed = time.monotonic() - started
    failures = dict(SOURCE_FAILURE_REASONS)
    if worker_failures:
        failures.update(worker_failures)
    LAST_EXTERNAL_STATUS = {
        "attempted": bool(jobs),
        "failed": bool(failures),
        "elapsed": elapsed,
        "results": len(unique),
        "source_counts": source_counts,
        "failures": failures,
    }
    terminal_log("WEB", f"completed in {elapsed:.1f}s | unique results={len(unique)}", "green" if unique else "yellow")
    return unique


def deduplicate_results(items: list[Result]) -> list[Result]:
    out = []
    seen = set()
    for r in items:
        key = (r.url or r.title).strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out



# ---------------------------------------------------------------------------
# Quick research on a brand-new topic
# ---------------------------------------------------------------------------

QUICK_SOURCE_FAMILIES = (
    "google", "web", "openalex", "semantic_scholar", "arxiv", "crossref",
    "dblp", "youtube", "github", "reddit", "quora",
)


def make_quick_query_plan(question: str, conversation_context: str = "") -> dict[str, Any]:
    prompt = f"""
You are a fast research search strategist.

New research topic:
{question}

Conversation memory for continuity (not evidence):
{conversation_context[:20000]}

This is NOT a deep literature review. The goal is to answer the question quickly using a
small, high-signal set of fresh sources. Choose up to {QUICK_MAX_SOURCE_FAMILIES} source
families from this list:
{', '.join(QUICK_SOURCE_FAMILIES)}

Use:
- academic sources for scholarly claims,
- Google/web for current or broad information,
- GitHub for code/implementation questions,
- Reddit/Quora for practitioner/community experiences,
- YouTube for demonstrations/tutorials,
- DBLP for computer-science coverage.

Return ONLY JSON with:
{{
  "source_families": ["..."],
  "queries": {{"source_family": ["1-2 precise queries"]}}
}}
Do not choose a source family just for completeness. Prefer the smallest useful set.
"""
    text = invoke_model(QUICK_QUERY_MODEL, prompt, ctx=QUICK_QUERY_CTX,
                        tokens_out=QUICK_QUERY_TOKENS, heartbeat=ANSWER_HEARTBEAT)
    try:
        m = re.search(r"\{.*\}", text, flags=re.S)
        data = json.loads(m.group(0)) if m else {}
        chosen = [x for x in data.get("source_families", []) if x in QUICK_SOURCE_FAMILIES]
        chosen = list(dict.fromkeys(chosen))[:QUICK_MAX_SOURCE_FAMILIES]
        qmap = data.get("queries", {}) if isinstance(data.get("queries", {}), dict) else {}
        queries = {fam: list(dict.fromkeys(qmap.get(fam, [])))[:2] for fam in chosen}
        if not chosen:
            raise ValueError("No source families selected")
        return {"source_families": chosen, "queries": queries}
    except Exception:
        fallback = ["google", "openalex", "arxiv", "github", "reddit"]
        return {
            "source_families": fallback,
            "queries": {fam: [question] for fam in fallback},
        }


def run_quick_searches(plan: dict[str, Any]) -> list[Result]:
    funcs = {
        "google": google_search,
        "web": web_search,
        "openalex": openalex_search,
        "semantic_scholar": semantic_scholar_search,
        "arxiv": arxiv_search,
        "crossref": crossref_search,
        "dblp": dblp_search,
        "youtube": youtube_search,
        "github": github_search,
        "reddit": reddit_search,
        "quora": quora_search,
    }
    jobs = []
    for family in plan.get("source_families", []):
        fn = funcs.get(family)
        if fn is None or not source_ready(family):
            continue
        for q in plan.get("queries", {}).get(family, [])[:2]:
            jobs.append((family, fn, q))

    def execute(job):
        family, fn, query = job
        try:
            vals = fn(query, QUICK_RESULTS_PER_SOURCE)
            for r in vals:
                r.metadata = dict(r.metadata or {})
                r.metadata["query"] = query
                r.metadata["quick_research"] = True
            return vals
        except Exception as exc:
            print(f"[{family}] quick research failed: {exc}")
            return []

    results: list[Result] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=QUICK_MAX_WORKERS) as pool:
        futures = [pool.submit(execute, job) for job in jobs]
        for fut in concurrent.futures.as_completed(futures):
            results.extend(fut.result())
    return deduplicate_results(results)


def quick_context(results: list[Result]) -> str:
    groups: dict[str, list[Result]] = {}
    for r in results:
        groups.setdefault(r.source, []).append(r)
    chunks = []
    for source, vals in groups.items():
        chunks.append(
            f"SOURCE {source.upper()}\n" + "\n".join(
                f"[{i+1}] {r.title}\nURL: {r.url}\nDATE: {r.date}\n{r.snippet[:1100]}\nMETA: {json.dumps(r.metadata or {}, ensure_ascii=False)}"
                for i, r in enumerate(vals[:QUICK_RESULTS_PER_SOURCE])
            )
        )
    return "\n\n".join(chunks)[:90000]


def answer_quick_question(question: str, plan: dict[str, Any], results: list[Result], conversation_context: str = "",
                          documents: list[Result] | None = None, document_project: str | None = None,
                          mode: str = MODE_AUTO) -> str:
    external_context=quick_context(results)
    doc_context=compact_context({},[],documents)
    combined=(doc_context+"\n\n" if doc_context else "")+external_context
    return answer_with_context(question,mode=mode,conversation_context=conversation_context,context=combined,
                               document_project=document_project)


def run_quick_research(question: str) -> int:
    print("\n=== QUICK RESEARCH: new topic ===")
    print("This is a targeted fresh evidence scan, not a deep literature review.\n")
    plan = make_quick_query_plan(question)
    print(f"[QUICK] source families: {', '.join(plan['source_families'])}")
    results = run_quick_searches(plan)
    print(f"[QUICK] collected {len(results)} unique results")
    answer = answer_quick_question(question, plan, results)

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = QUICK_RESEARCH_DIR / stamp
    out.mkdir(parents=True, exist_ok=True)
    (out / "question.txt").write_text(question + "\n", encoding="utf-8")
    (out / "query_plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "sources.json").write_text(json.dumps([asdict(x) for x in results], ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "answer.md").write_text(answer, encoding="utf-8")
    summary = {
        "mode": "quick",
        "status": "completed",
        "question": question,
        "title": question,
        "created": datetime.now(timezone.utc).isoformat(),
        "source_families": plan["source_families"],
        "source_count": len(results),
        "quick_research_dir": str(out),
    }
    (out / "run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n" + "=" * 100)
    print(answer)
    print("=" * 100)
    print(f"Quick research trace saved to: {out}")
    return 0

# ---------------------------------------------------------------------------
# Multimodal / vision helpers
# ---------------------------------------------------------------------------

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}

def validate_image_path(raw_path: str) -> Path:
    path = Path(os.path.expanduser(raw_path.strip().strip('\"\''))).resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Image not found: {path}")
    if path.suffix.lower() not in IMAGE_EXTENSIONS:
        raise ValueError(f"Unsupported image type: {path.suffix}. Supported: {', '.join(sorted(IMAGE_EXTENSIONS))}")
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > VISION_MAX_IMAGE_MB:
        raise ValueError(f"Image is {size_mb:.1f} MB; limit is {VISION_MAX_IMAGE_MB} MB.")
    try:
        from PIL import Image
        with Image.open(path) as img:
            img.verify()
    except Exception as exc:
        raise ValueError(f"Could not validate image: {exc}") from exc
    return path


def _available_image_files() -> list[Path]:
    files = []
    if not IMAGE_DIR.exists():
        return files
    for path in sorted(IMAGE_DIR.rglob("*"), key=lambda p: str(p).lower()):
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
            files.append(path)
    return files


def _resolve_image_input(raw: str) -> Path:
    cleaned = raw.strip().strip("\"'")
    if cleaned.isdigit():
        files = _available_image_files()
        idx = int(cleaned)
        if idx < 1 or idx > len(files):
            raise ValueError(f"Image number {idx} is not available.")
        return files[idx - 1]
    path = Path(os.path.expanduser(cleaned))
    if not path.is_absolute():
        # Prefer the standard workspace image folder; also accept an explicit
        # relative path such as images/ROV/plot.png.
        candidate = IMAGE_DIR / path
        if candidate.exists():
            return candidate
        workspace_candidate = ROOT / path
        if workspace_candidate.exists():
            return workspace_candidate
    return path


def attach_image_to_state(state: ConversationState) -> None:
    if not VISION_ENABLED:
        print("Image input is disabled in configuration.")
        return
    if len(state.attached_images or []) >= VISION_MAX_IMAGES:
        print(f"Maximum of {VISION_MAX_IMAGES} attached images reached. Clear images first with 'clear-image'.")
        return

    files = _available_image_files()
    print(f"Image folder: {IMAGE_DIR}")
    if files:
        print("Available images:")
        for idx, path in enumerate(files, 1):
            rel = path.relative_to(IMAGE_DIR)
            print(f"  {idx}. {rel}")
        print("Enter a number, a filename/path, or an absolute path.")
    else:
        print("No images found in the images/ folder. You can still enter an absolute path.")

    raw = input("Image: ").strip()
    if not raw:
        return
    try:
        path = validate_image_path(str(_resolve_image_input(raw)))
        state.attached_images = list(state.attached_images or [])
        resolved = str(path)
        if resolved in state.attached_images:
            print(colour(f"[IMAGE] already attached: {path.name}", "yellow"))
            return
        state.attached_images.append(resolved)
        save_conversation_state(state)
        print(colour(f"[IMAGE] attached: {path.name}", "green"))
    except Exception as exc:
        print(colour(f"[IMAGE] attach failed: {exc}", "red"))


def clear_attached_images(state: ConversationState) -> None:
    count = len(state.attached_images or [])
    state.attached_images = []
    save_conversation_state(state)
    print(colour(f"[IMAGE] cleared {count} attachment(s)", "yellow"))


def analyze_attached_images(question: str, image_paths: list[str]) -> str:
    if not VISION_ENABLED or not image_paths:
        return ""

    images = []
    seen_paths = set()
    for raw in image_paths[:VISION_MAX_IMAGES]:
        path = Path(raw)
        try:
            path = validate_image_path(str(path))
            if str(path) in seen_paths:
                continue
            seen_paths.add(str(path))
        except Exception as exc:
            print(colour(f"[IMAGE] skipped {raw}: {exc}", "yellow"))
            continue
        mime = mimetypes.guess_type(path.name)[0] or "image/jpeg"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        images.append((path, mime, encoded))

    if not images:
        return ""

    prompt = f"""
You are the dedicated visual-evidence stage of a local research assistant.
Analyze the attached image(s) specifically to provide reliable evidence for a
separate reasoning model that will answer the user's question.

User question:
{question}

Return a concise VISUAL EVIDENCE REPORT using these sections when applicable:
1. OBSERVED — directly visible text, objects, components, values, connections,
   equations, plots, tables, labels, and spatial relationships.
2. INFERRED — reasonable interpretations that go beyond direct observation.
3. UNCERTAIN / NOT VISIBLE — details that cannot be reliably determined.

Instructions:
- Read visible labels, equations, axis names, legends, table entries, diagram
  annotations, component names, and other text when legible.
- For plots and diagrams, describe important visible trends and relationships.
- For engineering/robotics images, identify visible components and spatial
  relationships only when reasonably clear.
- Preserve numerical values exactly when readable; do not guess unreadable values.
- Do not invent missing values, hidden content, or context outside the image.
- Do not answer the broader user question yet; provide evidence for the main
  language model to reason over.
""".strip()

    messages = [{"role": "user", "content": prompt, "images": [item[2] for item in images]}]
    started = time.monotonic()
    terminal_log("VISION", f"{VISION_MODEL} | analyzing {len(images)} image(s)...", "blue")
    payload = {
        "model": VISION_MODEL,
        "messages": messages,
        "stream": False,
        # Release the vision model after each analysis so the main answer
        # model can use the GPU/CPU memory on the next stage.
        "keep_alive": 0 if VISION_SERIAL else "5m",
        "options": {"temperature": 0, "num_ctx": VISION_CTX, "num_predict": VISION_TOKENS},
    }
    try:
        r = requests.post(f"{OLLAMA_BASE_URL}/api/chat", json=payload, timeout=VISION_TIMEOUT_SECONDS)
        r.raise_for_status()
        data = r.json()
        content = ((data.get("message") or {}).get("content")) or data.get("response") or ""
        content = str(content).strip()
        if not content:
            raise RuntimeError("Vision model returned an empty response")
        elapsed = time.monotonic() - started
        terminal_log("VISION", f"completed in {elapsed:.1f}s", "green")
        labels = []
        for idx, (path, _, _) in enumerate(images, 1):
            labels.append(f"[IMAGE-{idx}] {path.name}")
        return "ATTACHED IMAGE(S) VISUAL EVIDENCE\n" + "\n".join(labels) + "\n\n" + content
    except Exception as exc:
        terminal_log("VISION", f"analysis failed: {exc}", "red")
        return f"Visual analysis failed for the attached image(s): {exc}"


# ---------------------------------------------------------------------------
# Answer generation
# ---------------------------------------------------------------------------


def compact_context(local: dict[str, list[Result]], external: list[Result], documents: list[Result] | None = None) -> str:
    chunks=[]
    if documents:
        chunks.append("PERSONAL DOCUMENTS\n" + "\n".join(
            f"[DOC-{i+1}] {r.title}\nPATH: {r.url}\nPAGE: {r.date}\n{r.snippet[:3000]}\nMETA: {json.dumps(r.metadata or {}, ensure_ascii=False)}"
            for i,r in enumerate(documents)
        ))
    for family,vals in local.items():
        if vals:
            chunks.append(f"LOCAL RESEARCH {family.upper()}\n"+"\n".join(
                f"[LOCAL-{family}-{i+1}] {r.title}\nURL: {r.url}\n{r.snippet[:1200]}\nMETA: {json.dumps(r.metadata or {}, ensure_ascii=False)}"
                for i,r in enumerate(vals)))
    groups={}
    for r in external: groups.setdefault(r.source,[]).append(r)
    for source,vals in groups.items():
        chunks.append(f"EXTERNAL {source.upper()}\n"+"\n".join(
            f"[{source.upper()}-{i+1}] {r.title}\nURL: {r.url}\nDATE: {r.date}\n{r.snippet[:1000]}\nMETA: {json.dumps(r.metadata or {}, ensure_ascii=False)}"
            for i,r in enumerate(vals[:EXTERNAL_RESULTS_PER_SOURCE])))
    return "\n\n".join(chunks)[:FOLLOWUP_CONTEXT_CHARS]


def answer_with_context(question: str, *, mode: str, conversation_context: str,
                        context: str, research_run: dict[str,Any] | None = None,
                        document_project: str | None = None, web_only: bool = False,
                        web_notice: str = "", visual_context: str = "") -> str:
    run_text = ""
    document_text = ""
    if research_run and str(mode).upper().startswith("RESEARCH"):
        run_text = f"Research project: {research_run['title']}\\nOriginal question: {research_run['question']}\\nResearch date: {research_run['created'][:10]}\\nAge: {research_run['age_days']} days\\nRefresh: {research_run['refresh']}"
    if document_project and str(mode).upper().startswith("DOCUMENT"):
        document_text = f"Document project: {document_project}"
    prompt=f"""
You are a rigorous local research assistant.

Current mode: {mode}
{run_text}
{document_text}

User question:
{question}

Conversation memory (continuity only; NOT evidence):
{conversation_context[:18000]}

Source policy for this answer:
- CHAT: answer from the local LLM and conversation memory. Do not imply external verification.
- CHAT + WEB: answer from the local LLM plus the retrieved fresh web results; distinguish web evidence from reasoning.
- RESEARCH [DATABASE]: use only the selected/saved research corpus plus conversation memory.
- RESEARCH [WEB]: use fresh research/web evidence plus conversation memory.
- RESEARCH [DATABASE + WEB]: use the saved research corpus and fresh research/web evidence, keeping them distinguishable.
- RESEARCH [LLM ONLY]: use conversation continuity and model reasoning without retrieved research/web evidence.
- DOCUMENTS_ONLY: use only the user's indexed documents. Do not use outside knowledge to fill gaps.
- DOCUMENTS + WEB: use the user's indexed documents plus fresh web results, keeping them distinguishable.

Rules:
- Never invent facts or citations.
- Keep the answer focused on the user's actual question; avoid unnecessary exposition.
- Treat conversation memory as continuity, not evidence.
- Do not claim a document says something unless the retrieved text supports it.
- For technical/scientific claims, prefer primary or scholarly evidence when available.
- Clearly identify statements based on the user's documents versus research/web sources.
- When DOCUMENTS_ONLY is active and the documents do not support the answer, say that directly.
- Cite document evidence as [DOC-1], [DOC-2], etc. with filename and page.
- Cite external/local research using their tags only when such evidence is actually present.
- For CHAT with no web retrieval, do NOT cite web pages, do NOT invent sources, and do NOT include a Sources used section.
- For any mode with no retrieved evidence at all, do NOT fabricate citations or source titles.
- Include a compact Sources used section only when retrieved evidence is actually present.
- If a web-access note is provided below, state it plainly and do not imply that fresh web verification was available.
- When visual evidence is present, treat it as model-derived observation from the attached image(s).
- Do not turn uncertain visual interpretations into facts.
- When useful, refer to visual evidence as [IMAGE-1], [IMAGE-2], etc.

Web-access note for this question:
{web_notice}

Visual evidence from the dedicated vision model:
{visual_context}

Retrieved evidence:
{context}
"""
    print(f"[ANSWER] model={ANSWER_MODEL} context={len(context):,} chars | ctx={ANSWER_CTX} | output_cap={ANSWER_TOKENS}", flush=True)
    answer = invoke_model(ANSWER_MODEL,prompt,ctx=ANSWER_CTX,tokens_out=ANSWER_TOKENS,heartbeat=ANSWER_HEARTBEAT)
    # Hard guard against hallucinated web/source sections in pure LLM-only answers.
    if not context.strip() and not visual_context.strip():
        import re as _re
        answer = _re.split(r"\n\s*Sources used:\s*", answer, maxsplit=1, flags=_re.IGNORECASE)[0].rstrip()
    if web_notice:
        note = f"\n\n[WEB ACCESS NOTE] {web_notice}"
        if note.strip() not in answer:
            answer = answer.rstrip() + note
    return answer


def answer_question(question: str, run: dict[str, Any], local: dict[str, list[Result]], external: list[Result],
                    conversation_context: str = "", documents: list[Result] | None = None,
                    mode: str = MODE_RESEARCH_ONLY, document_project: str | None = None, visual_context: str = "") -> str:
    context=compact_context(local,external,documents)
    return answer_with_context(question,mode=mode,conversation_context=conversation_context,context=context,
                               research_run=run,document_project=document_project,visual_context=visual_context)


# ---------------------------------------------------------------------------
# Main interactive flow
# ---------------------------------------------------------------------------


def save_followup(run: dict[str, Any], question: str, query_plan: dict[str, Any],
                  local: dict[str, list[Result]], external: list[Result], answer: str,
                  conversation_state: ConversationState | None = None) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = FOLLOWUP_RUNS_DIR / stamp
    out.mkdir(parents=True, exist_ok=True)
    (out / "question.txt").write_text(question + "\n", encoding="utf-8")
    (out / "query_plan.json").write_text(json.dumps(query_plan, ensure_ascii=False, indent=2), encoding="utf-8")
    serial_local = {k: [asdict(x) for x in v] for k, v in local.items()}
    (out / "local_evidence.json").write_text(json.dumps(serial_local, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "external_sources.json").write_text(json.dumps([asdict(x) for x in external], ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "answer.md").write_text(answer, encoding="utf-8")
    if conversation_state is not None:
        (out / "conversation_state.json").write_text(
            json.dumps(asdict(conversation_state), ensure_ascii=False, indent=2), encoding="utf-8"
        )
    meta = {
        "created": datetime.now(timezone.utc).isoformat(),
        "research_project": run,
        "question": question,
        "external_sources_count": len(external),
        "source_status": {k: {"disabled": SOURCE_DISABLED[k], "errors": SOURCE_ERRORS[k]} for k in SOURCE_CONFIG},
    }
    (out / "followup_summary.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def reset_source_state() -> None:
    """Reset per-question source health so one failed question does not poison the next."""
    global LAST_EXTERNAL_STATUS
    SOURCE_FAILURE_REASONS.clear()
    LAST_EXTERNAL_STATUS = {"attempted": False, "failed": False, "elapsed": 0.0, "results": 0, "source_counts": {}, "failures": {}}
    for name in SOURCE_CONFIG:
        SOURCE_ERRORS[name] = 0
        SOURCE_DISABLED[name] = False
        SOURCE_LAST[name] = 0.0


def build_self_contained_deep_question(question: str, state: ConversationState) -> str:
    mem=conversation_prompt_context(state)
    if not mem.strip():
        return question
    prompt=f"""
Rewrite the user's current request into ONE self-contained deep-research question.
Preserve the exact scientific/technical intent. Resolve references using the conversation memory.
Do not add new goals. Do not answer the question.

Current request:
{question}

Conversation memory:
{mem[:28000]}

Return only the self-contained question.
"""
    try:
        return invoke_model(QUICK_QUERY_MODEL,prompt,ctx=QUICK_QUERY_CTX,tokens_out=900,heartbeat=ANSWER_HEARTBEAT).strip() or question
    except Exception:
        return question


def _latest_matching_research_run(records: list[dict[str, Any]], question: str) -> dict[str, Any] | None:
    """Return the newest completed run whose question/title is related to the handoff."""
    if not records:
        return None
    q = re.sub(r"\s+", " ", question.strip().lower())
    q_tokens = {t for t in re.findall(r"[a-z0-9]{4,}", q) if t not in STOP}
    candidates = []
    for rec in records:
        text = f"{rec.get('title','')} {rec.get('question','')}".lower()
        score = sum(1 for tok in q_tokens if tok in text)
        candidates.append((score, rec.get("created", ""), rec))
    candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return candidates[0][2] if candidates else records[0]


def run_deep_research_same_terminal(question: str, state: ConversationState | None = None) -> int:
    """Launch Deep Research Engine and optionally return to this assistant conversation."""
    engine_rel = DEEP_RESEARCH_ENGINE_SCRIPT
    engine = Path(engine_rel) if Path(engine_rel).is_absolute() else (ROOT / engine_rel)
    if not engine.exists():
        print(f"Deep research engine not found: {engine}")
        return 1

    final_question=build_self_contained_deep_question(question,state) if state else question
    print("\n[DEEP] Starting Deep Research Engine")
    print(f"[DEEP] Engine    : {engine}")
    print(f"[DEEP] Mode      : {DEEP_RESEARCH_ENGINE_MODE if DEEP_RESEARCH_PASS_MODE else 'engine default'}")
    print(f"[DEEP] Question  : {final_question}\n")

    cmd=[sys.executable,str(engine)]
    if DEEP_RESEARCH_PASS_MODE:
        cmd += ["--mode", DEEP_RESEARCH_ENGINE_MODE]
    cmd += [final_question]
    workdir = Path(DEEP_RESEARCH_WORKING_DIR).expanduser() if DEEP_RESEARCH_WORKING_DIR else engine.parent
    try:
        proc=subprocess.run(cmd,cwd=str(workdir),check=False)
        print(f"\n[DEEP] Deep Research Engine exited with code {proc.returncode}.")
        if proc.returncode != 0 or not DEEP_RESEARCH_AUTO_RESUME:
            return int(proc.returncode)

        # Rebuild the registry so the new run becomes visible immediately.
        try:
            new_records = build_registry()
        except Exception as exc:
            print(f"[DEEP] Could not refresh research registry: {exc}")
            return int(proc.returncode)

        new_run = _latest_matching_research_run(new_records, final_question)
        if state is not None and new_run is not None:
            sync_active_memory_scope(state)
            state.research_run = new_run
            state.project = new_run
            state.topic = new_run.get("title") or state.topic
            state.intention = INTENT_RESEARCH
            state.research_use_database = True
            state.research_web_search = True
            # Start/resume the research-specific memory scope without importing
            # the Chat/Document conversation history.
            activate_memory_scope(state, preserve_current=False)
            state.sync_legacy_mode()
            save_conversation_state(state)
            print(f"[DEEP] Attached completed run: {new_run.get('title','(untitled)')}")
            print("[DEEP] Returning to the Research Assistant conversation.")
        return int(proc.returncode)
    except KeyboardInterrupt:
        print("\n[DEEP] Deep research was interrupted.")
        return 130


def save_quick_conversation(question: str, plan: dict[str, Any], results: list[Result], answer: str,
                            project_label: str = "new-topic",
                            conversation_state: ConversationState | None = None) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = QUICK_RESEARCH_DIR / stamp
    out.mkdir(parents=True, exist_ok=True)
    (out / "question.txt").write_text(question + "\n", encoding="utf-8")
    (out / "query_plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "sources.json").write_text(json.dumps([asdict(x) for x in results], ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "answer.md").write_text(answer, encoding="utf-8")
    if conversation_state is not None:
        (out / "conversation_state.json").write_text(
            json.dumps(asdict(conversation_state), ensure_ascii=False, indent=2), encoding="utf-8"
        )
    summary = {
        "mode": "quick",
        "status": "completed",
        "title": question,
        "question": question,
        "created": datetime.now(timezone.utc).isoformat(),
        "project_label": project_label,
        "source_families": plan.get("source_families", []),
        "source_count": len(results),
    }
    (out / "run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def _enabled_intents() -> list[str]:
    out=[]
    if CHAT_ENABLED: out.append(INTENT_CHAT)
    if RESEARCH_ENABLED: out.append(INTENT_RESEARCH)
    if DOCUMENTS_ENABLED: out.append(INTENT_DOCUMENTS)
    return out


def intention_menu(current: str | None = None) -> str:
    options = _enabled_intents()
    if not options:
        raise RuntimeError("All assistant categories are disabled in .env.research_assistant")
    print("\nWhat do you want to do?")
    labels = {
        INTENT_CHAT: "Normal chat",
        INTENT_RESEARCH: "Research",
        INTENT_DOCUMENTS: "Document understanding",
    }
    for i,name in enumerate(options,1):
        mark=" *" if name==current else ""
        print(f"  {i}. {labels[name]}{mark}")
    while True:
        choice=input("Select intention: ").strip()
        try:
            idx=int(choice)-1
            if 0<=idx<len(options): return options[idx]
        except ValueError: pass
        print("Invalid selection.")


def yes_no_menu(prompt: str, default: bool) -> bool:
    default_label = "Y/n" if default else "y/N"
    while True:
        raw=input(f"{prompt} [{default_label}]: ").strip().lower()
        if not raw:
            return default
        if raw in {"y","yes","1","true","on"}: return True
        if raw in {"n","no","0","false","off"}: return False
        print("Please enter y or n.")


def chat_configuration(state: ConversationState) -> None:
    if CHAT_ALLOW_WEB_SEARCH:
        state.chat_web_search = yes_no_menu("Use LLM + web search?", state.chat_web_search)
    else:
        state.chat_web_search = False
    state.sync_legacy_mode()


def web_enabled_for_state(state: ConversationState) -> bool:
    if state.intention == INTENT_CHAT:
        return state.chat_web_search
    if state.intention == INTENT_RESEARCH:
        return state.research_web_search
    if state.intention == INTENT_DOCUMENTS:
        return state.documents_web_search
    return False

def set_web_for_state(state: ConversationState, enabled: bool) -> None:
    if state.intention == INTENT_CHAT:
        state.chat_web_search = enabled
    elif state.intention == INTENT_RESEARCH:
        state.research_web_search = enabled
    elif state.intention == INTENT_DOCUMENTS:
        state.documents_web_search = enabled
    state.sync_legacy_mode()

def toggle_web_for_state(state: ConversationState) -> None:
    if state.intention == INTENT_CHAT and not CHAT_ALLOW_WEB_SEARCH:
        terminal_log("WEB", "web search is disabled by configuration for CHAT", "yellow")
        return
    if state.intention == INTENT_RESEARCH and not RESEARCH_ALLOW_WEB_SEARCH:
        terminal_log("WEB", "web search is disabled by configuration for RESEARCH", "yellow")
        return
    if state.intention == INTENT_DOCUMENTS and not DOCUMENTS_ALLOW_WEB_SEARCH:
        terminal_log("WEB", "web search is disabled by configuration for DOCUMENTS", "yellow")
        return
    enabled = not web_enabled_for_state(state)
    set_web_for_state(state, enabled)
    save_conversation_state(state)
    terminal_log("WEB", f"web search {'ON' if enabled else 'OFF'} for subsequent questions", "green" if enabled else "yellow")


def research_configuration(state: ConversationState, records: list[dict[str, Any]]) -> tuple[dict[str, Any] | None, str | None]:
    # Choose an existing deep-research project or a genuinely new topic.
    print("\nResearch source")
    print("  1. Existing completed research project")
    print("  2. New research topic")
    while True:
        choice=input("Select: ").strip()
        if choice == "1":
            run=choose_run(records) if records else None
            if run is None:
                print("No completed research projects are available; choose a new topic.")
                continue
            topic=None
            break
        if choice == "2":
            topic=input("New research topic/question: ").strip()
            if topic:
                run=None
                break
        print("Invalid selection.")

    if RESEARCH_ALLOW_DATABASE:
        state.research_use_database=yes_no_menu(
            "Include the saved research database?",
            RESEARCH_DEFAULT_DATABASE,
        )
    else:
        state.research_use_database=False

    if RESEARCH_ALLOW_WEB_SEARCH:
        state.research_web_search=yes_no_menu(
            "Keep fresh web search active?",
            RESEARCH_DEFAULT_WEB_SEARCH,
        )
    else:
        state.research_web_search=False

    state.research_run=run
    state.project=run
    state.sync_legacy_mode()
    return run, topic


def document_configuration(state: ConversationState) -> None:
    state.document_project=choose_document_project(state.document_project)
    if DOCUMENTS_REQUIRE_PROJECT and not state.document_project:
        raise RuntimeError("Document understanding requires a document project.")
    if DOCUMENTS_ALLOW_WEB_SEARCH:
        state.documents_web_search=yes_no_menu(
            "Also use web search?",
            DOCUMENTS_DEFAULT_WEB_SEARCH,
        )
    else:
        state.documents_web_search=False
    state.sync_legacy_mode()


def print_session_header(state: ConversationState) -> None:
    print("\n"+"="*100)
    print(colour(f"Intention    : {state.intention}", "cyan"))
    if state.intention == INTENT_RESEARCH:
        print(colour(f"Research     : {state.research_run['title'] if state.research_run else '(none / new topic)'}", "cyan"))
        print("Doc project  : (inactive)")
    elif state.intention == INTENT_DOCUMENTS:
        print("Research     : (inactive)")
        print(colour(f"Doc project  : {state.document_project or '(none)'}", "cyan"))
    else:
        print("Research     : (inactive)")
        print("Doc project  : (inactive)")
    if state.intention == INTENT_CHAT:
        print(colour(f"Web search   : {'ON' if state.chat_web_search else 'OFF'}", "green" if state.chat_web_search else "yellow"))
    elif state.intention == INTENT_RESEARCH:
        print(f"Database     : {'ON' if state.research_use_database else 'OFF'}")
        print(colour(f"Web search   : {'ON' if state.research_web_search else 'OFF'}", "green" if state.research_web_search else "yellow"))
    else:
        print(colour(f"Web search   : {'ON' if state.documents_web_search else 'OFF'}", "green" if state.documents_web_search else "yellow"))
    print(f"Memory turns : {len(state.turns or [])} recent | summary={'yes' if state.summary.strip() else 'no'}")
    image_state = f"{len(state.attached_images or [])} attached" if state.attached_images else "none"
    print(f"Images       : {image_state} | vision={VISION_MODEL if VISION_ENABLED else 'OFF'}")
    print(f"Image folder : {IMAGE_DIR}")
    print("Commands     : i=intention | img=attach image | clear-image=remove images | w=toggle web | p=project | docs=document manager | sync=scan docs | deep=deep research | n=new topic | q=quit")
    print("="*100)


def run_web_only(question: str) -> tuple[list[Result], dict[str,Any], dict[str,Any]]:
    started = time.monotonic()
    plan={"web_queries":[question],"academic_queries":[],"youtube_queries":[],"github_queries":[],"reddit_queries":[],"quora_queries":[],"technical_queries":[question]}
    terminal_log("WEB", "fresh web search requested for this question", "blue")
    google_values = google_search(question, EXTERNAL_RESULTS_PER_SOURCE)
    web_values = web_search(question, EXTERNAL_RESULTS_PER_SOURCE)
    vals = deduplicate_results(google_values + web_values)
    failures = dict(SOURCE_FAILURE_REASONS)
    status={
        "attempted": True,
        "failed": bool(failures) and not vals,
        "partial_failure": bool(failures) and bool(vals),
        "elapsed": time.monotonic()-started,
        "results": len(vals),
        "failures": failures,
    }
    if status["failed"]:
        terminal_log("WEB", f"web access unavailable for this question after {status['elapsed']:.1f}s", "red")
    else:
        terminal_log("WEB", f"completed in {status['elapsed']:.1f}s | results={len(vals)}", "green" if vals else "yellow")
    return vals, plan, status


def global_research_results(question: str) -> dict[str, list[Result]]:
    """Search across saved completed research runs when no single project is selected."""
    records = build_registry()
    scored: dict[str, list[tuple[float, Result]]] = {"papers": [], "evidence": [], "technical": []}
    for run in records[:max(1, RESEARCH_DATABASE_MAX_PROJECTS)]:
        try:
            local = local_results(run, question)
            for family, values in local.items():
                for result in values:
                    blob = f"{result.title} {result.snippet}"
                    scored[family].append((score_text(question, blob), result))
        except Exception:
            continue
    out={}
    for family, values in scored.items():
        values.sort(key=lambda x:x[0], reverse=True)
        out[family]=[r for _,r in values[:LOCAL_PAPER_RESULTS if family=='papers' else LOCAL_EVIDENCE_RESULTS if family=='evidence' else LOCAL_TECH_RESULTS]]
    return out


def interactive_topic_session(run: dict[str, Any] | None, initial_question: str | None = None,
                               *, new_topic: bool = False, initial_intention: str | None = None,
                               initial_document_project: str | None = None,
                               initial_chat_web: bool | None = None,
                               initial_research_database: bool | None = None,
                               initial_research_web: bool | None = None,
                               initial_documents_web: bool | None = None) -> int:
    label=run["title"] if run else (initial_question or "New Topic")
    state=load_latest_conversation_state(label,run,initial_document_project)
    records=build_registry()
    if state is None:
        state=ConversationState(topic=label,intention=initial_intention or INTENT_CHAT,research_run=run,
                                document_project=initial_document_project,session_id=datetime.now().strftime("%Y%m%d_%H%M%S"))
        if initial_chat_web is not None: state.chat_web_search=initial_chat_web
        if initial_research_database is not None: state.research_use_database=initial_research_database
        if initial_research_web is not None: state.research_web_search=initial_research_web
        if initial_documents_web is not None: state.documents_web_search=initial_documents_web
    else:
        if initial_intention: state.intention=initial_intention
        if initial_document_project is not None: state.document_project=initial_document_project
        if initial_chat_web is not None: state.chat_web_search=initial_chat_web
        if initial_research_database is not None: state.research_use_database=initial_research_database
        if initial_research_web is not None: state.research_web_search=initial_research_web
        if initial_documents_web is not None: state.documents_web_search=initial_documents_web
        print(f"\n[CONVERSATION] Resuming {state.session_id} with {len(state.turns or [])} retained turns.")
    state.project=state.research_run
    state.sync_legacy_mode()
    save_conversation_state(state)

    if state.document_project:
        try: scan_project_documents(state.document_project)
        except Exception as exc: print(f"[RAG] Initial scan failed: {exc}")

    question=(initial_question or "").strip()
    while True:
        print_session_header(state)
        if not question:
            try: question=input(colour("\nQuestion: ", "cyan")).strip()
            except (EOFError,KeyboardInterrupt):
                save_conversation_state(state); return 0
            if not question: continue

        if question.lower() in {"i","intent","intention"}:
            try:
                sync_active_memory_scope(state)
                chosen=intention_menu(state.intention)
                state.intention=chosen
                if chosen==INTENT_CHAT:
                    chat_configuration(state)
                elif chosen==INTENT_RESEARCH:
                    run,topic=research_configuration(state,records)
                    state.research_run=run; state.project=run
                    if run:
                        state.topic=run.get("title") or state.topic
                    elif topic:
                        state.topic=topic
                        question=topic
                else:
                    document_configuration(state)
                # Load only the memory belonging to the newly selected scope.
                activate_memory_scope(state, preserve_current=False)
                save_conversation_state(state)
                if chosen != INTENT_RESEARCH or not question:
                    question=""
                continue
            except Exception as exc:
                print(f"[CONFIGURATION] {exc}")
                question=""; continue

        if question.lower() in {"p","project"}:
            sync_active_memory_scope(state)
            if state.intention==INTENT_RESEARCH:
                run=choose_run(records) if records else None
                if run:
                    state.research_run=run; state.project=run; state.topic=run["title"]
            elif state.intention==INTENT_DOCUMENTS:
                state.document_project=choose_document_project(state.document_project)
            else:
                state.document_project=choose_document_project(state.document_project)
            activate_memory_scope(state, preserve_current=False)
            state.sync_legacy_mode(); save_conversation_state(state); question=""; continue
        if question.lower() in {"docs","documents"}:
            state.document_project=document_manager(state.document_project); save_conversation_state(state); question=""; continue
        if question.lower() in {"sync","scan"}:
            if state.document_project:
                print(f"[RAG] {state.document_project}: {scan_project_documents(state.document_project)}")
            else:
                for name in document_projects(): print(f"[RAG] {name}: {scan_project_documents(name)}")
            question=""; continue
        if question.lower() in {"img", "attach", "image"}:
            attach_image_to_state(state)
            question = ""
            continue
        if question.lower() in {"clear-image", "clear-images", "rimg"}:
            clear_attached_images(state)
            question = ""
            continue
        if question.lower() in {"deep","d"}:
            sync_active_memory_scope(state)
            save_conversation_state(state)
            code = run_deep_research_same_terminal(
                question="Please conduct deep research on the current discussion.",
                state=state,
            )
            if code != 0:
                print(f"[DEEP] Research engine returned code {code}.")
                question = ""
                continue
            # The completed deep run has been attached to state; remain in this session.
            records = build_registry()
            question = ""
            continue
        if question.lower() in {"w", "web"}:
            toggle_web_for_state(state)
            question=""
            continue
        if question.lower()=="n":
            sync_active_memory_scope(state)
            topic=input(f"New topic/question ({state.intention}): ").strip()
            if topic:
                # Start a genuinely new topic while preserving the current intention
                # and its web/database/document settings. Do not re-run startup defaults.
                # For Research, a new topic is not tied to the previously selected run.
                if state.intention == INTENT_RESEARCH:
                    state.research_run = None
                    state.project = None
                state.summary = ""
                state.turns = []
                state.attached_images = []
                state.topic = topic
                new_scope = memory_scope_key(state)
                state.memory_scopes[new_scope] = {"summary": "", "turns": []}
                state.sync_legacy_mode()
                save_conversation_state(state)
                # Keep the topic captured above as the next query. The main loop will
                # execute it immediately instead of prompting for the same question again.
                terminal_log("SESSION", f"new topic started; web remains {'ON' if web_enabled_for_state(state) else 'OFF'}", "cyan")
                # Use the topic entered after `n` as the next question immediately.
                question = topic
            else:
                question=""
            continue
        if question.lower()=="q":
            save_conversation_state(state); return 0

        reset_source_state()
        question_started=time.monotonic()
        terminal_log("QUESTION", question, "cyan")
        conversation_context=conversation_prompt_context(state)
        documents=[]
        visual_context = ""
        if state.attached_images and VISION_ENABLED:
            visual_context = analyze_attached_images(question, list(state.attached_images))
        if state.intention==INTENT_DOCUMENTS and state.document_project:
            try:
                documents=rag_search(state.document_project,question,RAG_TOP_K)
                source_counts={}
                for doc in documents:
                    source_counts[doc.title]=source_counts.get(doc.title,0)+1
                print(f"[RAG] {len(documents)} document chunks retrieved from {state.document_project}")
                if RAG_VERBOSE_LOGGING and source_counts:
                    print("[RAG] Source documents:")
                    grouped={}
                    for doc in documents:
                        key=doc.url or doc.title
                        grouped.setdefault(key,{"title":doc.title.split(" — ",1)[0],"pages":set(),"chunks":0})
                        grouped[key]["chunks"]+=1
                        page=str((doc.metadata or {}).get("page") or doc.date or "").strip()
                        if page:
                            grouped[key]["pages"].add(page)
                    for item in sorted(grouped.values(),key=lambda x:(-x["chunks"],x["title"])):
                        pages=", ".join(sorted(item["pages"],key=lambda x:int(x) if x.isdigit() else 10**9))
                        page_text=f" | pages {pages}" if pages else ""
                        print(f"      {item['chunks']:>2} chunk(s) | {item['title']}{page_text}")
            except Exception as exc:
                print(f"[RAG] Retrieval failed: {exc}")

        local={"papers":[],"evidence":[],"technical":[]}
        external=[]
        plan={}
        try:
            web_notice=""
            if state.intention==INTENT_CHAT:
                if state.chat_web_search:
                    external,plan,web_status=run_web_only(question)
                    if web_status.get("failed"):
                        web_notice="Fresh web access could not be reached for this question, so I answered without fresh web evidence. The persistent Web search setting remains ON for the next question."

            elif state.intention==INTENT_DOCUMENTS:
                if state.documents_web_search:
                    external,plan,web_status=run_web_only(question)
                    if web_status.get("failed"):
                        web_notice="Fresh web access could not be reached for this question, so I answered using the available document context without fresh web evidence. The persistent Web search setting remains ON for the next question."

            elif state.intention==INTENT_RESEARCH:
                if state.research_use_database:
                    if state.research_run:
                        local=local_results(state.research_run,question)
                    else:
                        local=global_research_results(question)
                if state.research_web_search:
                    if state.research_run:
                        plan=make_query_plan(question,state.research_run,conversation_context)
                    else:
                        plan=make_quick_query_plan(question,conversation_context)
                    external=run_external_searches(plan)
                    if LAST_EXTERNAL_STATUS.get("failed") and not external:
                        web_notice="Fresh web access could not be reached for this question, so I answered from the available research database/local context without fresh web evidence. The persistent Web search setting remains ON for the next question."

            terminal_log("RETRIEVAL", f"intention={state.intention} images={len(state.attached_images or [])} documents={len(documents)} local={sum(len(v) for v in local.values())} external={len(external)}", "blue")
            if external:
                counts = LAST_EXTERNAL_STATUS.get("source_counts", {})
                if counts:
                    print(colour("[SOURCES] " + ", ".join(f"{k}={v}" for k,v in sorted(counts.items())), "dim"))
            if state.intention==INTENT_DOCUMENTS and DOCUMENTS_REQUIRE_PROJECT and not documents and not state.documents_web_search:
                answer=(f"I could not retrieve supporting content from the indexed documents in project "
                        f"'{state.document_project}'. Document understanding is configured without web search, "
                        f"so I will not fill the gap from outside knowledge.")
            else:
                context=compact_context(local,external,documents)
                # Effective mode applies only to this question. A failed web search
                # must not permanently change the user's persistent Web setting.
                if web_notice:
                    if state.intention==INTENT_CHAT:
                        retrieval_label="CHAT [WEB FALLBACK: OFF FOR THIS QUESTION]"
                    elif state.intention==INTENT_RESEARCH:
                        if state.research_use_database:
                            retrieval_label="RESEARCH [DATABASE | WEB FALLBACK: OFF FOR THIS QUESTION]"
                        else:
                            retrieval_label="RESEARCH [LLM ONLY | WEB FALLBACK: OFF FOR THIS QUESTION]"
                    else:
                        retrieval_label="DOCUMENTS_ONLY [WEB FALLBACK: OFF FOR THIS QUESTION]"
                else:
                    retrieval_label=state.mode
                answer=answer_with_context(question,mode=retrieval_label,conversation_context=conversation_context,context=context,
                                           research_run=state.research_run,document_project=state.document_project,
                                           web_notice=web_notice, visual_context=visual_context)
            add_conversation_turn(state,question,answer)
            refresh_conversation_summary(state)
            save_conversation_state(state)

            stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
            out=(FOLLOWUP_RUNS_DIR if state.research_run else QUICK_RESEARCH_DIR)/stamp
            out.mkdir(parents=True,exist_ok=True)
            (out/"question.txt").write_text(question+"\n",encoding="utf-8")
            (out/"answer.md").write_text(answer,encoding="utf-8")
            (out/"mode.json").write_text(json.dumps({
                "intention":state.intention,
                "mode":state.mode,
                "chat_web_search":state.chat_web_search,
                "research_use_database":state.research_use_database,
                "research_web_search":state.research_web_search,
                "documents_web_search":state.documents_web_search,
                "document_project":state.document_project,
                "research_run":state.research_run,
                "vision_enabled": VISION_ENABLED,
                "vision_model": VISION_MODEL,
                "attached_images": list(state.attached_images or [])
            },ensure_ascii=False,indent=2),encoding="utf-8")
            (out/"conversation_state.json").write_text(json.dumps(asdict(state),ensure_ascii=False,indent=2),encoding="utf-8")
            (out/"documents.json").write_text(json.dumps([asdict(x) for x in documents],ensure_ascii=False,indent=2),encoding="utf-8")
            (out/"attached_images.json").write_text(json.dumps(list(state.attached_images or []),ensure_ascii=False,indent=2),encoding="utf-8")
            
            if VISION_SAVE_EVIDENCE:
                (out/"visual_analysis.md").write_text(visual_context,encoding="utf-8")
            (out/"local_evidence.json").write_text(json.dumps({k:[asdict(x) for x in v] for k,v in local.items()},ensure_ascii=False,indent=2),encoding="utf-8")
            (out/"external_sources.json").write_text(json.dumps([asdict(x) for x in external],ensure_ascii=False,indent=2),encoding="utf-8")
            (out/"query_plan.json").write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding="utf-8")
            total_elapsed=time.monotonic()-question_started
            print("\n"+"="*100)
            print(colour(answer, "green"))
            print("="*100)
            terminal_log("TIMING", f"total question time = {total_elapsed:.1f}s", "magenta")
            if web_enabled_for_state(state):
                terminal_log("STATE", "persistent web setting remains ON", "green")
            print(f"Trace saved to: {out}")
        except KeyboardInterrupt:
            print("\n[INTERRUPTED] Search/inference interrupted.")
            save_conversation_state(state)
            try:
                choice=input("Run deep research for the current discussion? [y/N]: ").strip().lower()
            except (EOFError,KeyboardInterrupt):
                return 0
            if choice in {"y","yes"}:
                code = run_deep_research_same_terminal("Please conduct deep research on the current discussion.",state)
                if code != 0:
                    print(f"[DEEP] Research engine returned code {code}.")
                else:
                    records = build_registry()
                question = ""
                continue
        except Exception as exc:
            print(f"[ASSISTANT ERROR] {exc}")
        question=""

def parse_args() -> argparse.Namespace:
    p=argparse.ArgumentParser(description="Project-aware local assistant with three configurable user-intention categories")
    p.add_argument("question",nargs="*",help="Optional initial question")
    p.add_argument("--project",default=None,help="Deep-research run id/title substring")
    p.add_argument("--doc-project",default=None,help="Document-library project name")
    p.add_argument("--intention",choices=INTENTS,default=None,help="Initial user intention")
    p.add_argument("--mode",default=None,help="Legacy source mode; mapped to the new intention system")
    p.add_argument("--quick",action="store_true",help="Start a new research topic")
    p.add_argument("--no-interactive",action="store_true",help="Use the selected/new configuration without the startup selector")
    p.add_argument("--local-only",action="store_true",help="Legacy flag: Research intention with saved database and no web search")
    p.add_argument("--deep-mode",choices=("short","deep"),default=None,help="Override Deep Research Engine mode for this assistant session")
    p.add_argument("--self-test",action="store_true",help="Check config/RAG/Ollama connectivity without running research")
    return p.parse_args()


def legacy_to_intention(mode: str | None) -> str | None:
    if not mode:
        return None
    value=mode.strip().upper()
    if value == "DOCUMENTS_ONLY": return INTENT_DOCUMENTS
    if value in {"RESEARCH_ONLY", "DOCUMENTS + RESEARCH", "AUTO"}: return INTENT_RESEARCH
    if value == "WEB_SEARCH": return INTENT_CHAT
    if value == "CHAT": return INTENT_CHAT
    return None


def configure_startup(records: list[dict[str, Any]], args: argparse.Namespace) -> tuple[dict[str, Any] | None, str | None, dict[str, Any]]:
    """Run the three-category intention wizard and return run/topic/settings."""
    forced=args.intention or legacy_to_intention(args.mode)
    if forced in INTENTS and (args.intention or args.mode or args.no_interactive):
        intention=forced
    elif ASK_INTENTION_ON_START:
        intention=intention_menu()
    else:
        enabled=_enabled_intents()
        if not enabled:
            raise RuntimeError("All assistant categories are disabled")
        intention=enabled[0]

    settings={
        "intention": intention,
        "chat_web_search": CHAT_DEFAULT_WEB_SEARCH,
        "research_use_database": RESEARCH_DEFAULT_DATABASE,
        "research_web_search": RESEARCH_DEFAULT_WEB_SEARCH,
        "documents_web_search": DOCUMENTS_DEFAULT_WEB_SEARCH,
        "document_project": args.doc_project,
    }

    if intention==INTENT_CHAT:
        temp=ConversationState(topic="New Chat",intention=INTENT_CHAT)
        if args.mode and args.mode.strip().upper()=="WEB_SEARCH":
            temp.chat_web_search=True
        elif CHAT_ALLOW_WEB_SEARCH:
            temp.chat_web_search=yes_no_menu("Use LLM + web search?",CHAT_DEFAULT_WEB_SEARCH)
        else:
            temp.chat_web_search=False
        settings["chat_web_search"]=temp.chat_web_search
        return None,None,settings

    if intention==INTENT_RESEARCH:
        temp=ConversationState(topic="New Research",intention=INTENT_RESEARCH)
        if args.project:
            run=choose_run(records,args.project) if records else None
            topic=None
            if run is None:
                raise RuntimeError("Requested research project was not found")
        elif args.no_interactive and records:
            run=records[0]
            topic=None
        else:
            run,topic=research_configuration(temp,records)
            settings["research_use_database"]=temp.research_use_database
            settings["research_web_search"]=temp.research_web_search
            return run,topic,settings

        if args.local_only:
            temp.research_use_database=True
            temp.research_web_search=False
        else:
            temp.research_use_database=(yes_no_menu("Include the saved research database?",RESEARCH_DEFAULT_DATABASE)
                                        if RESEARCH_ALLOW_DATABASE else False)
            temp.research_web_search=(yes_no_menu("Keep fresh web search active?",RESEARCH_DEFAULT_WEB_SEARCH)
                                      if RESEARCH_ALLOW_WEB_SEARCH else False)
        settings["research_use_database"]=temp.research_use_database
        settings["research_web_search"]=temp.research_web_search
        return run,topic,settings

    # Document understanding
    temp=ConversationState(topic="Document Understanding",intention=INTENT_DOCUMENTS,document_project=args.doc_project)
    if not temp.document_project:
        temp.document_project=choose_document_project(None)
    if DOCUMENTS_REQUIRE_PROJECT and not temp.document_project:
        raise RuntimeError("Document understanding requires a document project")
    if DOCUMENTS_ALLOW_WEB_SEARCH:
        temp.documents_web_search=yes_no_menu("Also use web search?",DOCUMENTS_DEFAULT_WEB_SEARCH)
    else:
        temp.documents_web_search=False
    settings["document_project"]=temp.document_project
    settings["documents_web_search"]=temp.documents_web_search
    return None,None,settings


def self_test() -> int:
    print(f"Assistant env : {ASSISTANT_ENV_PATH}")
    print(f"Document YAML : {DOCUMENT_CONFIG_PATH}")
    print(f"RAG database  : {RAG_DB_PATH}")
    print(f"Ollama        : {OLLAMA_BASE_URL}")
    print(f"Models        : {QUICK_QUERY_MODEL} / {ANSWER_MODEL}")
    print(f"RAG embeddings: {RAG_EMBEDDINGS_ENABLED} ({RAG_EMBEDDING_MODEL})")
    print(f"RAG verbose   : {RAG_VERBOSE_LOGGING}")
    print(f"Categories    : chat={CHAT_ENABLED}, research={RESEARCH_ENABLED}, documents={DOCUMENTS_ENABLED}")
    print(f"Deep engine   : {ROOT / DEEP_RESEARCH_ENGINE_SCRIPT}")
    print(f"Deep mode     : {DEEP_RESEARCH_ENGINE_MODE}")
    print(f"Auto-resume   : {DEEP_RESEARCH_AUTO_RESUME}")
    init_rag_db()
    try:
        r=requests.get(f"{OLLAMA_BASE_URL}/api/tags",timeout=10); r.raise_for_status(); print("Ollama: OK")
    except Exception as exc:
        print(f"Ollama: unavailable ({exc})")
    return 0


def main() -> int:
    global DEEP_RESEARCH_ENGINE_MODE
    args=parse_args()
    if args.deep_mode:
        DEEP_RESEARCH_ENGINE_MODE = args.deep_mode
    ensure_document_library(); init_rag_db()
    if RAG_INDEX_ON_START:
        for name in document_projects():
            try: print(f"[RAG] Startup scan {name}: {scan_project_documents(name)}")
            except Exception as exc: print(f"[RAG] Startup scan failed for {name}: {exc}")
    if args.self_test: return self_test()
    records=build_registry()

    try:
        run,topic,settings=configure_startup(records,args)
    except (EOFError,KeyboardInterrupt):
        print("\nExiting.")
        return 0
    except Exception as exc:
        print(f"[STARTUP] {exc}")
        return 2

    initial_question=" ".join(args.question).strip() if args.question else topic
    if run is None and settings.get("intention")==INTENT_CHAT and topic is None and initial_question:
        initial_question=initial_question

    if settings.get("intention")==INTENT_DOCUMENTS:
        initial_doc=settings.get("document_project") or args.doc_project
        return interactive_topic_session(None,initial_question,new_topic=False,initial_intention=INTENT_DOCUMENTS,
                                         initial_document_project=initial_doc,initial_documents_web=settings.get("documents_web_search"))

    if settings.get("intention")==INTENT_RESEARCH:
        return interactive_topic_session(run,initial_question,new_topic=(run is None),initial_intention=INTENT_RESEARCH,
                                         initial_research_database=settings.get("research_use_database"),
                                         initial_research_web=settings.get("research_web_search"))

    return interactive_topic_session(None,initial_question,new_topic=True,initial_intention=INTENT_CHAT,
                                     initial_chat_web=settings.get("chat_web_search"))


if __name__ == "__main__":
    raise SystemExit(main())
